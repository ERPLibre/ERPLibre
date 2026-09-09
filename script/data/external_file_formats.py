#!/usr/bin/env python3
# © 2021-2026 TechnoLibre (http://www.technolibre.ca)
# License AGPL-3.0 or later (http://www.gnu.org/licenses/agpl)

"""Les formats : lire, décrire, nettoyer, écrire, convertir.

Séparé d'`external_file.py`, qui porte les RÈGLES. La frontière n'est pas
esthétique : `external_file.py` doit s'importer sous `.venv.erplibre`, qui
n'a ni openpyxl ni xlrd, sinon les tests des règles pures tombent tous à
l'import. Ici, chaque bibliothèque tierce est importée DANS la fonction qui
en a besoin, pour la même raison.

Le contrat externe reste `external_file.py --report/--plan/--apply` : la
coupure est interne.
"""

from __future__ import annotations

import collections
import csv
import io
import json
import math
import os
import re
import sys
import tempfile
import warnings
import xml.etree.ElementTree as ET

sys.path.append(
    os.path.normpath(os.path.join(os.path.dirname(__file__), "..", ".."))
)

from script.data import external_file as noyau  # noqa: E402
from script.data.external_file import (  # noqa: E402
    _INTACTE,
    Correspondance,
    anonymise_cellule,
    cellule_en_portee,
    classer,
    coercer_texte,
    colonne_plancher,
    detect_format,
    format_divergent,
    has_macros,
    nom_de_fichier_sur,
    normaliser_access,
    normaliser_xls,
    progres,
    valeur_forme_identifiant,
    valeur_hors_tableur,
    verifier_copie,
    vivier_de_mots,
)

# openpyxl avertit sur stderr (« DrawingML support is incomplete »), et
# stderr est le canal de progression : un avertissement s'y lirait comme
# une ligne de progrès.
warnings.simplefilter("ignore")

CIBLES_CONVERSION = ("xlsx", "csv", "json", "xml")

# Le nom de feuille d'un format qui n'en porte pas. Reprendre le nom du
# FICHIER le recrachait dans la copie convertie — clé de premier niveau d'un
# JSON, nom d'onglet d'un xlsx — alors que le dialogue promet que le nom du
# fichier n'est pas anonymisé.
NOM_FEUILLE_NEUTRE = "feuille_1"

# `csv` refuse un champ de plus de 128 Kio par défaut, et un mémo de
# commande les dépasse. Le refus arrivait en « format non reconnu » sur un
# CSV parfaitement valide.
csv.field_size_limit(16 * 1024 * 1024)


class ErreurMoteur(Exception):
    """Un refus motivé, porteur d'une clé d'`ERREURS`.

    `conseil` est une SECONDE clé, celle du remède, quand le refus en a un
    que l'opérateur ne devinerait pas. Le détail nomme ce que le moteur
    trouve, le conseil dit quoi répondre à la prochaine exécution ; les
    mêler dans une chaîne les rend intraduisibles tous les deux.
    """

    def __init__(self, cle, detail="", conseil=""):
        super().__init__(cle)
        self.cle = cle
        self.detail = detail
        self.conseil = conseil


# ----------------------------------------------------------------------
# Ce que cette machine sait lire
# ----------------------------------------------------------------------
def _importable(nom):
    try:
        __import__(nom)
        return True
    except Exception:
        return False


def capabilities():
    """Format -> lisible ici. Les formats stdlib le sont toujours."""
    excel = _importable("openpyxl")
    return {
        "csv": True,
        "json": True,
        "xml": True,
        "xlsx": excel,
        "xls": _importable("xlrd"),
        "xlsb": False,
        "access": _importable("access_parser"),
        "images": _importable("PIL"),
        "mots": _importable("randomwordfr"),
    }


# ----------------------------------------------------------------------
# La grille — la forme commune à tous les lecteurs
# ----------------------------------------------------------------------
class Feuille:
    """Un tableau nommé : des lignes de valeurs, déjà normalisées.

    Chaque lecteur ramène ses types aux types des règles AVANT de rendre
    une feuille. C'est là que se joue la correction d'un `.xls` dont les
    dates sont des flottants ou d'un Access dont les dates sont des
    chaînes : la règle ne doit jamais avoir à deviner.
    """

    def __init__(self, nom, lignes, masquee=False, source=None):
        self.nom = nom
        self.lignes = lignes
        self.masquee = masquee
        self.source = source
        # (ligne, colonne) -> (conteneur, clé). Hors tableur, l'écriture
        # passait par un SECOND parcours de l'arbre, qui ignorait la portée,
        # le plancher et les colonnes intactes — au point de détruire un
        # external ID que le plancher venait de protéger. Une ancre par
        # cellule rend l'écriture solidaire de la marche à blanc.
        self.ancres = {}
        # L'arbre JSON ou XML dont ces ancres sont les branches.
        self.arbre = None
        # L'index 1-based de la colonne qui porte de la STRUCTURE, ou None.
        # Le LECTEUR seul le sait : une grille (clé, valeur) en a une, un
        # tableau de vrais champs n'en a aucune. Le déduire du FORMAT
        # recopiait en clair le premier champ d'un tableau d'objets.
        self.colonne_structure = None
        # Vrai quand la ligne 1 de la grille est une ÉTIQUETTE fabriquée
        # par le lecteur (« cle »/« valeur », « chemin »/« valeur ») :
        # absente du fichier, elle n'est ni de la donnée ni un en-tête, et
        # aucune ancre ne peut l'écrire.
        self.ligne1_fabriquee = False
        # Le porteur d'une racine SCALAIRE (`"x"`, `42` — du JSON valide).
        # Un str ou un int ne se réécrit pas en place : sans cette case
        # réinscriptible, l'ancre porte None et l'écriture déréférence None.
        self.porteur = None
        # La ligne 1-based qui NOMME les colonnes, ou None quand la feuille
        # n'en a pas. Un tableur ne le dit pas : la ligne 1 est une
        # PRÉSOMPTION, fausse sur un rapport dont A1 porte un titre et sur
        # une feuille sans en-tête, où elle recopiait de la donnée en
        # clair. `_preparer` la mesure ; les lecteurs qui FABRIQUENT leur
        # ligne 1 — Access, JSON, XML — la connaissent d'avance.
        self.ligne_champs = 1
        # L'empan des lignes d'en-tête, ligne de champs comprise : ce qui
        # les surmonte sans être des données — un titre, une catégorie
        # fusionnée — reste intact avec elle.
        self.lignes_entete = {1}
        # Ce que le FICHIER DÉCLARE, quand il le déclare : `headerRowCount`
        # d'un tableau d'un onglet. La déclaration GAGNE sur la mesure —
        # `_resynchroniser_tableaux` la lit déjà pour nommer les colonnes
        # d'un tableau, et deux notions d'en-tête qui se contredisent
        # feraient renommer un tableau depuis une ligne qu'on vient
        # d'anonymiser.
        self.entete_declaree = None
        # Les premières lignes telles que le lecteur les a LUES, avant
        # coercition. Un csv rend tout en chaînes ; les coercer donne au
        # signal de type la matière dont il a besoin, mais une ligne
        # d'en-tête doit garder ses chaînes — « 2024 » est un libellé, pas
        # un nombre. `mesurer_entetes` les rend à l'empan retenu.
        self.lignes_brutes = None

    @property
    def etiquettes(self):
        """Les libellés de colonne, pris sur la ligne de CHAMPS.

        Vide quand la feuille n'a pas d'en-tête : sans libellé, le
        plancher ne peut pas reconnaître un identifiant et la question des
        colonnes intactes répond par l'index. C'est le prix juste — les
        prendre sur une ligne de données faisait planchéier au hasard.
        """
        rang = self.ligne_champs
        if rang is None or rang > len(self.lignes):
            return []
        return list(self.lignes[rang - 1])


def corps(feuille):
    """Les lignes de DONNÉES : toutes celles qui ne sont pas de l'en-tête.

    Un seul endroit le dit, pour les deux graveurs qui écrivaient
    `lignes[1:]` : une seconde ligne d'en-tête y devenait un
    enregistrement, et sur une feuille SANS en-tête la première ligne de
    données devenait les noms de clé — en clair, et perdue comme donnée.

    L'exclusion se fait par APPARTENANCE à l'empan, non par position sous
    lui. Une ligne de données au-dessus de la ligne de champs en est
    exclue exprès — la mesure la reconnaît, la portée l'anonymise — et
    partir de la dernière ligne d'en-tête la faisait DISPARAÎTRE d'une
    conversion json ou xml, comptée comme remplacée puis absente du
    fichier. Un empan corrigé par l'opérateur n'est d'ailleurs pas tenu
    d'être contigu.
    """
    empan = feuille.lignes_entete or set()
    return [
        ligne
        for rang, ligne in enumerate(feuille.lignes, start=1)
        if rang not in empan
    ]


def _etiquettes_par_colonne(feuilles):
    """{(feuille, colonne 1-based): étiquette}, lu AVANT toute écriture.

    Avant, parce que répondre « oui » à l'anonymisation de la ligne
    d'en-tête remplacerait l'étiquette que la question des colonnes
    intactes cherche ensuite.
    """
    table = {}
    for feuille in feuilles:
        for index, etiquette in enumerate(feuille.etiquettes, start=1):
            if isinstance(etiquette, str) and etiquette.strip():
                table[(feuille.nom, index)] = etiquette.strip()
    return table


# ----------------------------------------------------------------------
# La ligne d'en-tête : mesurée, non présumée
# ----------------------------------------------------------------------
# Au-delà, on ne cherche plus : un tableau dont la ligne de champs vient
# plus bas qu'ici n'est pas un tableau mais un rapport mis en page, et
# l'opérateur corrige mieux que n'importe quelle mesure.
LIGNES_SONDEES = 10

# Jusqu'où l'écran MONTRE des lignes. La portée de l'opérateur n'est pas
# bornée par celle de la mesure : un rapport dont la ligne de champs vient
# en douzième position se corrige à la main, et le commentaire ci-dessus
# le promettait sans que rien ne l'offre — l'écran s'arrêtait à la
# dixième, donc la douzième était inatteignable.
#
# Les lignes au-delà de `LIGNES_SONDEES` arrivent SANS mesure : la
# calculer coûte un balayage de la feuille par ligne, et trente balayages
# sur un classeur de deux millions de cellules se voient. Une colonne de
# mesures vide y dit la vérité — la mesure n'est pas allée jusque-là.
LIGNES_MONTREES = 30

# Une ligne qui RESSEMBLE à ses données en est. Le seuil sépare deux
# nuages mesurés sur les huit tables d'une base réelle, prises une fois
# avec leur ligne d'en-tête et une fois sans : les en-têtes vont de 0,00 à
# 0,36, les lignes de données de 0,50 à 1,00. Le seuil penche vers le bas
# de l'écart, parce que les deux erreurs ne coûtent pas la même chose —
# manquer un en-tête l'anonymise et abîme la lecture de la copie, en
# inventer un le recopie en clair et fait SORTIR de la donnée.
ACCORD_DE_DONNEE = 0.40

# Ce qu'une ligne de champs doit tenir par ailleurs. Le remplissage écarte
# le titre seul en A1 ; la distinction écarte la ligne de catégorie, qui
# répète un même mot sur plusieurs colonnes.
REMPLISSAGE_MINIMAL = 0.5
DISTINCTION_MINIMALE = 0.99

# L'un OU l'autre suffit, et l'un des deux est REQUIS : ce sont les seules
# preuves POSITIVES qu'une ligne nomme ses colonnes. `accord` bas dit
# seulement « elle ne ressemble pas à ses données », ce qui n'est pas la
# même chose — s'en contenter fait prendre une première ligne de
# données pour un en-tête sur une table tout-texte, donc la recopie en
# clair.
#
# Le contraste est franc sur une table numérique et MUET sur une table
# tout-texte ; l'appartenance à la colonne y répond, mais seulement là où
# la colonne a un VOCABULAIRE. Aucune des deux ne parle : pas d'en-tête.
#
# Le seuil du contraste est bas parce que toute ligne de DONNÉES mesurée
# vaut zéro — une ligne au-dessus d'une colonne numérique y est numérique
# elle aussi. Il n'est pas nul pour qu'une seule valeur texte égarée dans
# une colonne de nombres ne suffise pas : il en faut une part.
CONTRASTE_MINIMAL = 0.15
HORS_COLONNE_MINIMAL = 0.9


# Les lecteurs qui FABRIQUENT leur ligne 1 : elle porte des noms de champ
# ou des chemins de clé absents du fichier, donc l'empan est connu d'avance
# et rien n'est à mesurer.
FORMATS_ENTETE_FABRIQUEE = ("access", "json", "xml")


def mesurer_entetes(feuilles, format_lu, corrections=None, memoire=None):
    """Poser, par feuille, l'empan d'en-tête et la ligne de champs.

    Cinq sources, dans cet ordre de priorité :

    1. La CORRECTION de l'opérateur pour ce passage-ci, qui ne se
       remesure jamais.
    2. Ce que la TABLE se rappelle : une correction faite sur un fichier
       du même lot. C'est une réponse d'opérateur elle aussi, donc elle
       passe avant ce que le fichier déclare.
    3. Ce que le FICHIER déclare — `headerRowCount` d'un tableau.
    4. Ce que le LECTEUR sait : Access, JSON et XML fabriquent leur
       ligne 1, gardent le défaut et ne mesurent rien.
    5. La mesure, pour un tableur sans tableau déclaré et pour un csv.

    La ligne de CHAMPS est la plus BASSE de l'empan : c'est la convention
    de `_resynchroniser_tableaux`, qui nomme les colonnes d'un tableau
    depuis la dernière ligne de `headerRowCount`. L'adopter plutôt que
    d'en inventer une seconde évite deux notions qui se contredisent.
    """
    corrections = corrections or {}
    memoire = memoire or {}
    for feuille in feuilles:
        demandee = corrections.get(feuille.nom)
        if demandee is None:
            demandee = memoire.get(feuille.nom)
        if demandee is not None:
            empan = {int(n) for n in demandee if int(n) >= 1}
        elif feuille.entete_declaree:
            empan = set(feuille.entete_declaree)
        elif format_lu in FORMATS_ENTETE_FABRIQUEE or feuille.ligne1_fabriquee:
            empan = {1}
        else:
            empan, _champs = lignes_entete(feuille.lignes)
        feuille.lignes_entete = empan
        feuille.ligne_champs = max(empan) if empan else None
        _rendre_les_brutes(feuille)


def _rendre_les_brutes(feuille):
    """Rendre à l'empan ses valeurs telles que le lecteur les a LUES.

    Un csv arrive tout en chaînes et se fait coercer pour que le signal de
    type ait de la matière ; une ligne d'en-tête, elle, doit garder ses
    chaînes — « 2024 » y est un libellé de colonne, pas un nombre. Sans
    ce retour, l'étiquette d'une colonne annuelle devenait un entier, et
    la réponse « 2024 » à la question des colonnes ne portait plus.
    """
    brutes = feuille.lignes_brutes
    if not brutes:
        return
    for numero in feuille.lignes_entete:
        if 1 <= numero <= len(brutes) and numero <= len(feuille.lignes):
            feuille.lignes[numero - 1] = list(brutes[numero - 1])


# Ce qu'une colonne montre d'elle-même. Trois suffisent à reconnaître un
# champ ; borner la longueur garde la ligne lisible et la charge utile
# petite.
EXEMPLES_PAR_COLONNE = 3
EXEMPLE_LONGUEUR = 40

# Au-delà, les exemples ne sont plus peuplés. `colonnes` elle-même n'est
# JAMAIS tronquée : les bornes, les formes et le plancher en dérivent
# colonne par colonne, et une liste coupée déplancherait en silence.
EXEMPLES_COLONNES_MAX = 500


def valeur_d_exemple(valeur):
    """Une valeur montrable : toujours une `str`, toujours bornée.

    Écarte les octets — `valeur_hors_tableur` tomberait sur son `str()`
    final et rendrait le `repr` d'un `bytes`, donc du binaire en clair —
    et les flottants non finis, que `allow_nan=False` refuserait à la
    sérialisation, c'est-à-dire APRÈS tout le travail du moteur, l'écran
    n'affichant alors qu'une trace.

    Un seul type dans le champ : ni le menu ni l'écran n'ont à brancher.
    Cette valeur ne passe JAMAIS par `t()`, qui chercherait une clé de
    traduction dans une donnée du client.
    """
    if isinstance(valeur, (bytes, bytearray)):
        return "<%d octets>" % len(valeur)
    if isinstance(valeur, float) and not math.isfinite(valeur):
        return ""
    # `or ""` écraserait le zéro et le faux, qui sont de VRAIES valeurs :
    # une colonne de montants nuls montrait des exemples vides.
    rendu = valeur_hors_tableur(valeur)
    texte = "" if rendu is None else str(rendu)
    if len(texte) > EXEMPLE_LONGUEUR:
        return texte[: EXEMPLE_LONGUEUR - 1] + "…"
    return texte


def _forme_de_valeur(valeur):
    """La signature de forme d'une valeur : chiffres en 9, lettres en a.

    Les répétitions sont écrasées, si bien que « 99999 » et « 999999 »
    sont une seule forme. Une ligne de données partage la forme de sa
    colonne — « ZK204817 » au-dessus de « ZK204818 » — et un nom de champ
    non. C'est ce qui tranche là où le type ne dit rien, les deux étant du
    texte.
    """
    if valeur is None:
        return ""
    texte = re.sub(r"[0-9]", "9", str(valeur).strip())
    texte = re.sub(r"[^\W\d_]", "a", texte)
    return re.sub(r"(.)\1+", r"\1+", texte)


def _accord_de_forme(lignes, rang):
    """Part des colonnes où la ligne partage la forme dominante du corps.

    Haut : la ligne ressemble à ses données, donc c'en est. Bas : elle s'en
    distingue, donc elle les nomme. Rend None quand il n'y a pas de corps
    sous la ligne — rien à comparer n'est pas un verdict.
    """
    largeur = max((len(l) for l in lignes), default=0)
    if rang > len(lignes) or not largeur:
        return None
    ligne = lignes[rang - 1]
    corps = lignes[rang:]
    if not corps:
        return None
    accords = compares = 0
    for index in range(largeur):
        valeurs = [
            l[index]
            for l in corps
            if index < len(l) and classer(l[index]) != "vide"
        ]
        cellule = ligne[index] if index < len(ligne) else None
        if not valeurs or classer(cellule) == "vide":
            continue
        compares += 1
        dominante = collections.Counter(
            _forme_de_valeur(v) for v in valeurs
        ).most_common(1)
        if dominante and _forme_de_valeur(cellule) == dominante[0][0]:
            accords += 1
    return accords / compares if compares else 0.0


def _signaux_entete(lignes, rang):
    """Quatre mesures de la ligne `rang` comme ligne de champs.

    `contraste` : elle est du texte au-dessus d'une colonne qui n'en porte
    pas. `hors_colonne` : sa valeur ne figure pas parmi celles de sa
    colonne — zéro collision mesurée sur 55 colonnes tout-texte réelles.
    `rempli` : elle couvre la largeur utile. `distinct` : ses libellés ne
    se répètent pas.
    """
    largeur = max((len(l) for l in lignes), default=0)
    if rang > len(lignes) or not largeur:
        return None
    ligne = lignes[rang - 1]
    corps = lignes[rang:]
    if not corps:
        return None
    contraste = hors_colonne = compares = colonnes_a_vocabulaire = 0
    for index in range(largeur):
        valeurs = [l[index] for l in corps if index < len(l)]
        familles = {classer(v) for v in valeurs} - {"vide"}
        cellule = ligne[index] if index < len(ligne) else None
        if not familles or classer(cellule) == "vide":
            continue
        compares += 1
        if classer(cellule) == "texte" and "texte" not in familles:
            contraste += 1
        non_vides = [v for v in valeurs if classer(v) != "vide"]
        vues = {str(v).strip().lower() for v in non_vides}
        # Ce signal n'a de pouvoir que sur une colonne à VOCABULAIRE, où
        # des valeurs se répètent. Dans une colonne tout-distincte — des
        # noms, des courriels, des numéros de pièce — une ligne de DONNÉES
        # est absente du reste de sa colonne exactement autant qu'un
        # libellé l'est : rapporté à la largeur, le signal y vaut 1,0 pour
        # n'importe quelle ligne. Les compter laisse `accord` seul sur une
        # table tout-texte, où une première ligne de données ponctuée
        # autrement que son corps passe alors pour un en-tête : recopiée
        # en clair dans la copie.
        if len(vues) < len(non_vides):
            colonnes_a_vocabulaire += 1
            if str(cellule).strip().lower() not in vues:
                hors_colonne += 1
    pleines = sum(
        1
        for index in range(largeur)
        if index < len(ligne) and classer(ligne[index]) != "vide"
    )
    etiquettes = [
        str(ligne[index]).strip().lower()
        for index in range(min(largeur, len(ligne)))
        if classer(ligne[index]) != "vide"
    ]
    return {
        "contraste": contraste / compares if compares else 0.0,
        # Son propre dénominateur : les colonnes où il porte quelque
        # chose. Aucune n'en a, il vaut 0,0 — c'est-à-dire « ce signal ne
        # dit rien ici », et non « la ligne est de la donnée ».
        "hors_colonne": (
            hors_colonne / colonnes_a_vocabulaire
            if colonnes_a_vocabulaire
            else 0.0
        ),
        "rempli": pleines / largeur,
        "distinct": (
            len(set(etiquettes)) / len(etiquettes) if etiquettes else 0.0
        ),
        "compares": compares,
    }


def _est_une_ligne_de_champs(signaux, accord):
    """Le verdict, seuils nommés à l'appui.

    Une limite à connaître : l'accord de forme est AVEUGLE quand le nom
    d'un champ et ses valeurs partagent une classe de caractères — « nom »
    au-dessus de mots, sur une grille étroite et tout-alphabétique. Rien
    de structurel ne les sépare alors ; seul un vocabulaire le ferait, et
    une liste de noms de champs connus est une classe OUVERTE. Le verdict
    y est donc « pas d'en-tête », ce qui anonymise la ligne — le côté sur
    lequel pencher — et l'opérateur corrige.

    Une classe de LONGUEUR ne lève pas l'aveuglement : « ZK204817 » et
    « MPQ204818 » doivent s'accorder alors que « etiquette » et
    « aboulie » doivent se distinguer, et aucune frontière ne fait les
    deux — mesuré sur cinq jeux de seuils.
    """
    if signaux is None or accord is None:
        return False
    return (
        signaux["rempli"] >= REMPLISSAGE_MINIMAL
        and signaux["distinct"] >= DISTINCTION_MINIMALE
        and accord <= ACCORD_DE_DONNEE
        and (
            signaux["contraste"] >= CONTRASTE_MINIMAL
            or signaux["hors_colonne"] >= HORS_COLONNE_MINIMAL
        )
    )


def _est_de_la_donnee(lignes, rang):
    """Cette ligne, au-dessus de la ligne de champs, est-elle une DONNÉE ?

    Deux conditions, et les deux sont nécessaires : elle ressemble à ses
    données par la forme, ET elle en remplit la largeur. La forme seule ne
    suffit pas — un titre seul en A1 partage la forme de la colonne de
    noms qu'il surmonte, et sans la seconde l'extension s'arrête dessus et
    le met en portée. Le remplissage seul ne suffit pas non plus : une
    ligne de catégorie couvre la largeur sans être des données.

    L'erreur va du bon côté quand elle se produit : une ligne prise pour
    des données est ANONYMISÉE, non recopiée.
    """
    accord = _accord_de_forme(lignes, rang)
    if accord is None or accord < 0.5:
        return False
    signaux = _signaux_entete(lignes, rang)
    return signaux is not None and signaux["rempli"] >= REMPLISSAGE_MINIMAL


def lignes_entete(lignes):
    """(empan des lignes d'en-tête, ligne de champs) — 1-based.

    La ligne de CHAMPS est celle que la mesure retient : c'est elle qui
    nomme les colonnes, donc celle dont `etiquettes` sort. L'EMPAN y ajoute
    les lignes du dessus qui ne ressemblent pas à des données — un titre de
    rapport, une ligne de catégorie fusionnée : les anonymiser n'apporte
    rien et rend la copie illisible.

    Rend `(set(), None)` quand aucune ligne ne mesure comme une ligne de
    champs. C'est un verdict, pas un échec : une feuille sans en-tête
    existe, et sa ligne 1 est de la DONNÉE — la présumer d'en-tête la
    recopiait en clair.
    """
    if not lignes:
        return set(), None
    for rang in range(1, min(LIGNES_SONDEES, len(lignes)) + 1):
        accord = _accord_de_forme(lignes, rang)
        if not _est_une_ligne_de_champs(_signaux_entete(lignes, rang), accord):
            continue
        empan = {rang}
        haut = rang - 1
        while haut >= 1 and not _est_de_la_donnee(lignes, haut):
            empan.add(haut)
            haut -= 1
        return empan, rang
    return set(), None


def _stats_colonnes(feuille):
    """Par colonne : étiquette, type dominant, remplies, distinctes, bornes.

    Le coût est nul — la passe visite déjà chaque cellule — et c'est elle
    qui produit les BORNES dont la règle du nombre a besoin, en même temps
    que les étiquettes que teste le plancher. Sans ce bloc, les questions
    sur les colonnes et sur la ligne d'en-tête sont impossibles à
    répondre : il faudrait ouvrir le fichier dans Excel d'abord, ce que
    cette entrée existe pour supprimer.
    """
    largeur = max((len(l) for l in feuille.lignes), default=0)
    colonnes = []
    etiquettes = feuille.etiquettes
    for index in range(largeur):
        familles = {}
        distinctes = set()
        remplies = 0
        forme = True
        forme_rel = True
        mini = maxi = None
        entiere = True
        exemples = []
        for numero, ligne in enumerate(feuille.lignes, start=1):
            # L'EMPAN, jamais le littéral « 1 » : `forme_identifiant` et
            # `forme_relation` se court-circuitent sur un seul faux, si
            # bien qu'une ligne d'en-tête laissée dans la mesure fait
            # lâcher le plancher sur `partner_id`, `key`, `model` et
            # `state` — exactement les colonnes pour lesquelles il existe.
            if numero in (feuille.lignes_entete or ()):
                continue
            valeur = ligne[index] if index < len(ligne) else None
            famille = classer(valeur)
            if famille == "vide":
                continue
            remplies += 1
            forme = forme and valeur_forme_identifiant(valeur)
            forme_rel = forme_rel and noyau.valeur_forme_relation(valeur)
            familles[famille] = familles.get(famille, 0) + 1
            if len(distinctes) < 10000:
                avant = len(distinctes)
                try:
                    distinctes.add(valeur)
                except TypeError:
                    distinctes.add(repr(valeur))
                # Une valeur INÉDITE et rien qu'elle : trois exemples
                # identiques ne montrent rien de la colonne. La valeur est
                # déjà en main et la passe est déjà payée.
                if (
                    len(distinctes) > avant
                    and len(exemples) < EXEMPLES_PAR_COLONNE
                    and index < EXEMPLES_COLONNES_MAX
                ):
                    montrable = valeur_d_exemple(valeur)
                    if montrable:
                        exemples.append(montrable)
            if (
                famille == "nombre"
                and not isinstance(valeur, bool)
                and math.isfinite(valeur)
            ):
                # NaN et l'infini ne bornent rien : `min` et `max` les
                # ignorent selon l'ORDRE des arguments, ce qui n'est pas
                # une garantie sur laquelle asseoir un intervalle de
                # tirage.
                mini = valeur if mini is None else min(mini, valeur)
                maxi = valeur if maxi is None else max(maxi, valeur)
                # Le verdict porte sur TOUTE la colonne : une seule
                # décimale quelque part la rend décimale.
                entiere = entiere and float(valeur).is_integer()
        etiquette = etiquettes[index] if index < len(etiquettes) else None
        dominant = (
            max(familles.items(), key=lambda kv: kv[1])[0]
            if familles
            else "vide"
        )
        colonnes.append(
            {
                "index": index + 1,
                "etiquette": (
                    str(etiquette).strip()
                    if isinstance(etiquette, str)
                    else None
                ),
                "type": dominant,
                "remplies": remplies,
                "distinctes": len(distinctes),
                # CE qui distingue deux colonnes sans libellé : mesuré, ni
                # le type, ni le compte, ni les bornes n'y suffisent.
                "exemples": exemples,
                "min": mini,
                "max": maxi,
                "entiere": entiere and mini is not None,
                "forme_identifiant": forme,
                "forme_relation": forme_rel,
                # Une SÉLECTION est un ensemble fermé et petit. Sans cette
                # borne, une colonne de provinces ou de créneaux nommés par
                # des personnes passait pour une sélection sur le seul fait
                # d'être en minuscules.
                "selection": (
                    len(distinctes) <= noyau.SELECTION_MAX_DISTINCTES
                ),
                "plancher": colonne_plancher(
                    etiquette,
                    forme,
                    forme_rel,
                    len(distinctes) <= noyau.SELECTION_MAX_DISTINCTES,
                ),
            }
        )
    return colonnes


def _formes_par_colonne(rapport, cle="forme_identifiant"):
    """{(feuille, colonne): son contenu a-t-il cette forme mesurée}."""
    return {
        (feuille["nom"], colonne["index"]): bool(colonne.get(cle))
        for feuille in rapport.get("feuilles", [])
        for colonne in feuille.get("colonnes", [])
    }


def _bornes_par_colonne(rapport):
    """{(feuille, colonne): (min, max, entiere)} depuis le rapport.

    Le troisième terme dit si la colonne ne porte QUE des entiers. Il ne
    se déduit pas du type Python d'une valeur : `.xls` ne stocke que des
    doubles, et son lecteur rend 100 en `100.0`.
    """
    bornes = {}
    for feuille in rapport.get("feuilles", []):
        for colonne in feuille.get("colonnes", []):
            if colonne.get("min") is not None:
                bornes[(feuille["nom"], colonne["index"])] = (
                    colonne["min"],
                    colonne["max"],
                    colonne.get("entiere"),
                )
    return bornes


# ----------------------------------------------------------------------
# Les lecteurs
# ----------------------------------------------------------------------
def _lire_xlsx(chemin, garder_vba=False):
    """Le classeur, et ses feuilles de calcul.

    `wb.worksheets` et JAMAIS `wb.sheetnames` : ce dernier mêle les
    feuilles graphiques, et `wb[nom]` rend alors un `Chartsheet`, qui n'a
    ni `max_row`, ni `max_column`, ni `iter_rows`.
    """
    from openpyxl import load_workbook

    try:
        classeur = load_workbook(
            chemin,
            data_only=False,
            keep_links=False,
            keep_vba=bool(garder_vba),
        )
    except ErreurMoteur:
        raise
    except Exception as exc:
        # La bibliothèque ne lit pas tout ce qu'Excel écrit : une feuille
        # graphique DÉPOURVUE de graphique fait lever `AttributeError`
        # dans son lecteur de relations, en 3.1.2. Ce refus nomme la
        # cause ; sans lui, un fichier qui s'ouvre dans Excel ressort en
        # « format non reconnu » suivi d'un message Python.
        raise ErreurMoteur(
            "lecture_impossible", f"{type(exc).__name__}: {exc}"
        )
    feuilles = []
    for onglet in classeur.worksheets:
        lignes = [
            [cellule.value for cellule in ligne]
            for ligne in onglet.iter_rows()
        ]
        feuille = Feuille(
            onglet.title,
            lignes,
            masquee=onglet.sheet_state != "visible",
            source=onglet,
        )
        feuille.entete_declaree = _entete_declaree(onglet)
        feuilles.append(feuille)
    return classeur, feuilles


def _entete_declaree(onglet):
    """L'empan d'en-tête que les TABLEAUX de l'onglet déclarent, ou None.

    Un tableau OOXML porte `headerRowCount` et la première ligne de son
    `ref` : l'empan s'en déduit sans rien mesurer. C'est la même donnée que
    `_resynchroniser_tableaux` lit pour nommer les colonnes, et l'adopter
    ici évite deux notions d'en-tête qui se contredisent — un tableau
    renommé depuis une ligne qu'on vient d'anonymiser.

    Plusieurs tableaux sur un onglet : l'empan est leur RÉUNION, chacun
    gardant sa propre ligne d'en-tête intacte.
    """
    empan = set()
    for tableau in (getattr(onglet, "tables", {}) or {}).values():
        ref = getattr(tableau, "ref", "") or ""
        if ":" not in ref and not ref:
            continue
        try:
            premiere = onglet[ref.split(":")[0]].row
        except (ValueError, KeyError, TypeError):
            continue
        hauteur = getattr(tableau, "headerRowCount", 1)
        try:
            hauteur = int(hauteur if hauteur is not None else 1)
        except (TypeError, ValueError):
            hauteur = 1
        if hauteur < 1:
            # Un tableau déclaré SANS ligne d'en-tête : sa première ligne
            # est de la donnée, et rien n'est à garder pour lui.
            continue
        empan.update(range(premiere, premiere + hauteur))
    return empan or None


def _lire_xls(chemin):
    import xlrd

    # `logfile` vaut sys.stdout par défaut, et xlrd y écrit dès qu'un
    # classeur n'a pas de CODEPAGE : la ligne se mêlait à l'unique objet
    # JSON de stdout, et l'appelant refusait un fichier lisible sans un mot
    # de diagnostic.
    try:
        classeur = xlrd.open_workbook(
            chemin, formatting_info=False, logfile=sys.stderr
        )
    except Exception as exc:
        # La bibliothèque échoue sur un fichier MALFORMÉ, et pas seulement
        # sur un format qu'elle ignore : un flux de classeur abîmé, un nom
        # défini dont la formule ne s'évalue pas. Sans ce refus, la sortie
        # annonçait « format non reconnu » — faux, le format est reconnu —
        # suivi d'un message Python.
        raise ErreurMoteur(
            "lecture_impossible", f"{type(exc).__name__}: {exc}"
        )
    feuilles = []
    for onglet in classeur.sheets():
        lignes = []
        for numero in range(onglet.nrows):
            lignes.append(
                [
                    normaliser_xls(
                        onglet.cell_type(numero, colonne),
                        onglet.cell_value(numero, colonne),
                        classeur.datemode,
                    )
                    for colonne in range(onglet.ncols)
                ]
            )
        feuilles.append(
            Feuille(onglet.name, lignes, masquee=onglet.visibility != 0)
        )
    return feuilles


def _lire_access(chemin):
    """Les tables d'un `.mdb`/`.accdb`, reconstruites ligne par ligne.

    Deux pièges d'`access-parser`, tous deux dans sa source :

    `_parse_catalog` ajoute `MSysObjects` AVANT le filtre des objets
    système, sans condition. Itérer `db.catalog` livrerait donc le
    catalogue d'objets d'Access — noms, propriétaires, et pour une table
    liée le chemin source. On écarte tout `MSys*`.

    `parse_table` rend un dictionnaire de colonnes de LONGUEURS INÉGALES :
    `_parse_row` sort tôt après avoir déjà alimenté les colonnes de
    longueur fixe. Reconstruire par `zip(*valeurs)` tronquerait à la plus
    courte et apparierait un nom avec les nombres d'une autre ligne —
    pire qu'une erreur, parce que la sortie a l'air correcte. On aligne
    donc sur la plus LONGUE, en comblant par `None`.
    """
    from access_parser import AccessParser

    try:
        base = AccessParser(chemin)
    except Exception as exc:
        raise ErreurMoteur(
            "lecture_impossible", f"{type(exc).__name__}: {exc}"
        )
    feuilles = []
    for nom in base.catalog:
        if str(nom).startswith("MSys"):
            continue
        try:
            colonnes = base.parse_table(nom)
        except Exception as exc:  # une table illisible n'arrête pas tout
            progres(f"{nom}: {type(exc).__name__}")
            continue
        etiquettes = list(colonnes.keys())
        # Le TYPE déclaré de chaque colonne, pour normaliser avant que
        # l'anonymiseur voie quoi que ce soit : `access-parser` rend une
        # date et un montant en CHAÎNE, et la règle du texte les prendrait
        # pour du texte du client.
        types = {}
        try:
            for col in base.get_table(nom).columns.values():
                types[col.col_name_str] = col.type
        except Exception as exc:  # pragma: no cover - table hors norme
            progres(f"{nom}: types indisponibles ({type(exc).__name__})")
        hauteur = max((len(v) for v in colonnes.values()), default=0)
        lignes = [etiquettes]
        for index in range(hauteur):
            lignes.append(
                [
                    normaliser_access(
                        (
                            colonnes[cle][index]
                            if index < len(colonnes[cle])
                            else None
                        ),
                        types.get(cle),
                    )
                    for cle in etiquettes
                ]
            )
        feuilles.append(Feuille(str(nom), lignes))
    return feuilles


def _encodage_csv(chemin):
    """(encodage, par quoi il a été décidé).

    Le BOM d'abord, `chardet` ensuite s'il est là, puis un décodage d'essai
    du FICHIER ENTIER — pas d'un préfixe, qui couperait une séquence
    multi-octets et ferait passer de l'UTF-8 valide pour du cp1252 — et
    `cp1252` en dernier, qui ne lève jamais. Le rapport dit lequel a
    répondu, pour qu'une supposition ne se lise pas comme une mesure.
    """
    with open(chemin, "rb") as fh:
        octets = fh.read()
    for bom, nom in (
        (b"\xef\xbb\xbf", "utf-8-sig"),
        (b"\xff\xfe\x00\x00", "utf-32"),
        (b"\x00\x00\xfe\xff", "utf-32"),
        (b"\xff\xfe", "utf-16"),
        (b"\xfe\xff", "utf-16"),
    ):
        if octets.startswith(bom):
            return nom, "bom"
    try:
        import chardet

        devine = chardet.detect(octets)
        if devine and devine.get("encoding") and devine["confidence"] > 0.8:
            return devine["encoding"], "chardet"
    except Exception:
        pass
    try:
        octets.decode("utf-8")
        return "utf-8", "repli"
    except UnicodeDecodeError:
        return "cp1252", "repli"


def _delimiteur_csv(texte):
    """(délimiteur, par quoi il a été décidé).

    `csv.Sniffer().sniff()` LÈVE sur un CSV à une seule colonne et sur un
    fichier en dents de scie — exactement les deux formes que le rapport
    promet de couvrir — et rend une LETTRE prise dans la donnée si on ne le
    borne pas. Son verdict est donc borné, rattrapé, puis VALIDÉ : est
    retenu le candidat qui donne le même nombre de champs sur les vingt
    premières lignes non vides.
    """
    candidats = [",", ";", "\t", "|"]
    echantillon = texte[:65536]
    try:
        devine = csv.Sniffer().sniff(echantillon, delimiters=",;\t|")
        if devine.delimiter in candidats:
            candidats.insert(
                0, candidats.pop(candidats.index(devine.delimiter))
            )
            mesure = "sniffer"
        else:
            mesure = "repli"
    except csv.Error:
        mesure = "repli"

    lignes = [l for l in echantillon.splitlines() if l.strip()][:20]
    for candidat in candidats:
        largeurs = {len(l.split(candidat)) for l in lignes}
        if len(largeurs) == 1 and largeurs != {1}:
            return candidat, mesure if candidat == candidats[0] else "mesure"
    return ",", "repli"


def _lire_csv(chemin):
    encodage, source_enc = _encodage_csv(chemin)
    with open(chemin, "r", encoding=encodage, errors="replace") as fh:
        texte = fh.read()
    delimiteur, source_del = _delimiteur_csv(texte)
    # `csv.reader` ne rend que des chaînes : un champ numérique doit
    # retrouver son type ici, sinon la colonne de montants est traitée
    # comme du texte et chaque montant devient un mot.
    # L'EN-TÊTE garde son type d'origine : la coercition le rendait en
    # entier, donc en étiquette absente, et « 1 » ne désignait plus la
    # colonne étiquetée « 1 » mais la première colonne.
    lignes = []
    brutes = []
    for numero, ligne in enumerate(
        csv.reader(io.StringIO(texte), delimiter=delimiteur), start=1
    ):
        # TOUTES les lignes sont coercées, ligne 1 comprise : le signal de
        # type compare une ligne aux VRAIS types de sa colonne, et lui
        # laisser ses chaînes faisait passer toute ligne 1 pour un en-tête
        # — donc en inventer un là où il n'y en a pas, ce qui recopie de
        # la donnée en clair. `mesurer_entetes` rend ensuite leurs chaînes
        # aux seules lignes de l'empan retenu, où « 2024 » est un libellé.
        # La fenêtre des brutes couvre la portée de l'ÉCRAN, non celle
        # de la mesure : une ligne d'en-tête désignée au-delà retrouvait
        # ses valeurs coercées, et « 2024 » y redevenait un nombre.
        if numero <= LIGNES_MONTREES:
            brutes.append(list(ligne))
        lignes.append([coercer_texte(champ) for champ in ligne])
    nom = NOM_FEUILLE_NEUTRE
    meta = {
        "encodage": encodage,
        "encodage_source": source_enc,
        "delimiteur": delimiteur,
        "delimiteur_source": source_del,
    }
    feuille = Feuille(nom, lignes)
    feuille.lignes_brutes = brutes
    return feuille, meta


# ----------------------------------------------------------------------
# Les rapports
# ----------------------------------------------------------------------
def _rapport_commun(chemin, format_lu):
    return {
        "chemin": chemin,
        "format": format_lu,
        "taille": os.path.getsize(chemin),
        "divergence": format_divergent(chemin, format_lu),
        "vivier": len(vivier_de_mots()),
        "vivier_complet": _importable("randomwordfr"),
        "lecture_seule": format_lu in noyau.FORMATS_LECTURE_SEULE,
        "feuilles": [],
        "comptes": {},
        "hors_cellules": {},
        "avertissements": [],
    }


def _comptes_de(feuilles):
    comptes = {}
    for feuille in feuilles:
        for numero, ligne in enumerate(feuille.lignes, start=1):
            for valeur in ligne:
                famille = classer(valeur)
                if famille == "vide":
                    continue
                comptes[famille] = comptes.get(famille, 0) + 1
    return comptes


def _feuilles_en_rapport(feuilles):
    resume = []
    for feuille in feuilles:
        colonnes = _stats_colonnes(feuille)
        formules = sum(
            1
            for ligne in feuille.lignes
            for valeur in ligne
            if classer(valeur) == "formule"
        )
        litteraux = sum(
            1
            for ligne in feuille.lignes
            for valeur in ligne
            if classer(valeur) == "formule" and '"' in str(valeur)
        )
        empan = sorted(feuille.lignes_entete or ())
        resume.append(
            {
                "nom": feuille.nom,
                # Une LISTE, non un set : `json.dump` refuse un set, et
                # ce rapport traverse un sous-processus.
                "lignes_entete": empan,
                "ligne_champs": feuille.ligne_champs,
                "entete_mesure": _mesure_de_la_ligne(feuille),
                "lignes_sondees": _lignes_sondees(feuille),
                "entete_declaree": bool(feuille.entete_declaree),
                # COMBIEN de colonnes n'ont pas d'exemple, non un
                # simple « oui ». Un booléen que personne ne lisait ne
                # disait pas plus qu'une case vide, et une colonne sans
                # exemple se lit comme une colonne vide — l'opérateur
                # laisse alors intacte, ou non, une colonne qu'il n'a pas
                # vue.
                "exemples_manquants": max(
                    0,
                    max((len(l) for l in feuille.lignes), default=0)
                    - EXEMPLES_COLONNES_MAX,
                ),
                "lignes": len(feuille.lignes),
                "colonnes_n": max((len(l) for l in feuille.lignes), default=0),
                "masquee": feuille.masquee,
                "formules": formules,
                "formules_litteral": litteraux,
                "colonnes": colonnes,
            }
        )
    return resume


def _lignes_sondees(feuille):
    """Les premières lignes, montrables, avec ce que la mesure en dit.

    Sans elles, un écran ne peut pas faire juger QUELLE ligne nomme les
    colonnes : le rapport dit son verdict mais pas la matière sur laquelle
    il porte, et contredire un verdict qu'on ne voit pas est un pari.

    Les mesures accompagnent chaque ligne, pas seulement celle retenue :
    c'est ainsi qu'on voit pourquoi la voisine a été écartée.
    """
    rendu = []
    for rang in range(1, min(LIGNES_MONTREES, len(feuille.lignes)) + 1):
        ligne = feuille.lignes[rang - 1]
        apercu = [
            valeur_d_exemple(v)
            for v in ligne[:EXEMPLES_PAR_COLONNE]
            if classer(v) != "vide"
        ]
        signaux = accord = None
        if rang <= LIGNES_SONDEES:
            signaux = _signaux_entete(feuille.lignes, rang)
            accord = _accord_de_forme(feuille.lignes, rang)
        mesure = None
        if signaux is not None and accord is not None:
            mesure = {
                "contraste": round(signaux["contraste"], 2),
                "hors_colonne": round(signaux["hors_colonne"], 2),
                "accord": round(accord, 2),
            }
        rendu.append(
            {
                "numero": rang,
                "apercu": apercu,
                "pleines": sum(1 for v in ligne if classer(v) != "vide"),
                "mesure": mesure,
            }
        )
    return rendu


def _mesure_de_la_ligne(feuille):
    """Les cinq mesures de la ligne de champs, ou None.

    Affichées pour qu'une INVENTION se lise avant d'être consentie : la
    mesure peut prendre une ligne de données pour un en-tête, et
    l'opérateur doit voir POURQUOI elle a tranché avant de la contredire.
    """
    rang = feuille.ligne_champs
    if rang is None:
        return None
    signaux = _signaux_entete(feuille.lignes, rang)
    accord = _accord_de_forme(feuille.lignes, rang)
    if signaux is None or accord is None:
        return None
    return {
        cle: round(valeur, 2)
        for cle, valeur in list(signaux.items()) + [("accord", accord)]
        if cle != "compares"
    }


def _hors_cellules_xlsx(classeur, chemin):
    """Ce qui porte du texte sans être une cellule.

    Les plages nommées se lisent dans les DEUX collections : openpyxl 3.1 a
    scindé l'espace de noms, et `ws.defined_names` porte les noms locaux à
    une feuille, invisibles à `wb.defined_names`.

    Les liens externes et les hyperliens de cellule se comptent SÉPARÉMENT :
    ce sont des objets différents, et un seul des deux tombe avec
    `keep_links=False`.
    """
    plages = list(classeur.defined_names.keys())
    hyperliens, cibles, commentaires, auteurs = 0, set(), 0, set()
    entetes, validations, conditionnelles, graphiques = 0, 0, 0, 0
    croises = 0
    for onglet in classeur.worksheets:
        plages.extend(onglet.defined_names.keys())
        graphiques += len(getattr(onglet, "_charts", []) or [])
        croises += len(getattr(onglet, "_pivots", []) or [])
        validations += len(
            getattr(onglet.data_validations, "dataValidation", []) or []
        )
        conditionnelles += sum(1 for _ in onglet.conditional_formatting)
        for entete in (
            onglet.oddHeader,
            onglet.evenHeader,
            onglet.firstHeader,
            onglet.oddFooter,
            onglet.evenFooter,
            onglet.firstFooter,
        ):
            for partie in (entete.left, entete.center, entete.right):
                if getattr(partie, "text", None):
                    entetes += 1
        for ligne in onglet.iter_rows():
            for cellule in ligne:
                if cellule.comment is not None:
                    commentaires += 1
                    if cellule.comment.author:
                        auteurs.add(cellule.comment.author)
                if cellule.hyperlink is not None:
                    hyperliens += 1
                    cible = getattr(cellule.hyperlink, "target", None)
                    if cible:
                        cibles.add(str(cible))
    proprietes = classeur.properties
    return {
        "createur": proprietes.creator,
        "modifie_par": proprietes.lastModifiedBy,
        "titre": proprietes.title,
        "mots_cles": proprietes.keywords,
        "proprietes_perso": len(
            list(getattr(classeur, "custom_doc_props", []) or [])
        ),
        "commentaires": commentaires,
        "auteurs_commentaires": len(auteurs),
        "hyperliens": hyperliens,
        "cibles_liens": sorted(cibles)[:20],
        "liens_externes": len(getattr(classeur, "_external_links", []) or []),
        "plages_nommees": sorted(set(plages))[:40],
        "entetes_pieds": entetes,
        "validations": validations,
        "conditionnelles": conditionnelles,
        "graphiques": graphiques,
        "croises": croises,
        "feuilles_graphiques": len(getattr(classeur, "chartsheets", []) or []),
        "images": noyau.compter_media(chemin),
        "macros": has_macros(chemin),
    }


def report(chemin):
    """Le rapport, avant toute question."""
    format_lu = detect_format(chemin)
    if format_lu == "protege":
        raise ErreurMoteur("protege", chemin)
    if format_lu == "xlsb":
        rapport = _rapport_commun(chemin, "xlsb")
        rapport["hors_cellules"] = {"macros": has_macros(chemin)}
        rapport["arret"] = noyau.ERREURS["illisible_ici"]
        return rapport
    if not format_lu:
        raise ErreurMoteur("format_inconnu", os.path.basename(chemin))

    rapport = _rapport_commun(chemin, format_lu)
    if format_lu == "xlsx":
        progres(os.path.basename(chemin))
        classeur, feuilles = _lire_xlsx(chemin)
        rapport["hors_cellules"] = _hors_cellules_xlsx(classeur, chemin)
    elif format_lu == "xls":
        feuilles = _lire_xls(chemin)
        rapport["avertissements"].append(
            "This format stores no formula readable here."
        )
    elif format_lu == "access":
        feuilles = _lire_access(chemin)
        rapport["avertissements"].append(
            "Saved queries are not readable in pure Python."
        )
    elif format_lu == "csv":
        feuille, meta = _lire_csv(chemin)
        feuilles = [feuille]
        rapport.update(meta)
    elif format_lu == "json":
        feuilles = _lire_json(chemin)[0]
    elif format_lu == "xml":
        feuilles = _lire_xml(chemin)[0]
    else:  # pragma: no cover - detect_format ne rend rien d'autre
        raise ErreurMoteur("format_inconnu", format_lu)

    # AVANT les statistiques : elles sautent l'empan d'en-tête, et une
    # ligne d'en-tête restée dans la mesure rendrait `forme_identifiant`
    # faux pour toute la colonne — le plancher lâcherait alors les
    # colonnes pour lesquelles il existe.
    mesurer_entetes(feuilles, format_lu)
    rapport["feuilles"] = _feuilles_en_rapport(feuilles)
    rapport["comptes"] = _comptes_de(feuilles)
    if not rapport["vivier_complet"]:
        rapport["avertissements"].append(
            "Only 20 fallback words are available: randomwordfr is missing."
        )
    return rapport


# ----------------------------------------------------------------------
# JSON et XML — hors tableur, une valeur est aussi un attribut
# ----------------------------------------------------------------------
def _lire_json(chemin):
    """Un JSON ramené à une grille, pour que le rapport parle.

    Les CLÉS sont de la structure et ne bougent pas ; seules les valeurs
    passent par les règles.
    """
    with open(chemin, "r", encoding="utf-8") as fh:
        arbre = json.load(fh)
    lignes = []
    ancres = {}
    colonne_structure = None
    ligne1_fabriquee = False
    porteur = None
    # Le tableau d'enregistrements exige que TOUTE entrée soit un objet :
    # une entrée nue (chaîne, nombre, null, liste) n'a pas de clés, et
    # `element.get` la faisait mourir. Un tableau mêlé retombe sur
    # l'aplatissement, qui ancre chaque feuille — la sauter écrirait
    # l'entrée en clair, non annoncée.
    if (
        isinstance(arbre, list)
        and arbre
        and all(isinstance(e, dict) for e in arbre)
    ):
        cles = []
        for element in arbre:
            for cle in element:
                if cle not in cles:
                    cles.append(cle)
        lignes.append(cles)
        for element in arbre:
            lignes.append([element.get(cle) for cle in cles])
            for index, cle in enumerate(cles, start=1):
                if cle in element:
                    ancres[(len(lignes), index)] = (element, cle)
    else:
        colonne_structure = 1
        ligne1_fabriquee = True
        lignes.append(["cle", "valeur"])
        porteur = None if isinstance(arbre, (dict, list)) else [arbre]
        for cle, valeur, conteneur, index in _aplatir_json(
            arbre,
            conteneur=porteur,
            index=None if porteur is None else 0,
        ):
            lignes.append([cle, valeur])
            ancres[(len(lignes), 2)] = (conteneur, index)
    feuille = Feuille(NOM_FEUILLE_NEUTRE, lignes)
    feuille.ancres = ancres
    feuille.arbre = arbre
    feuille.colonne_structure = colonne_structure
    feuille.ligne1_fabriquee = ligne1_fabriquee
    feuille.porteur = porteur
    return [feuille], arbre


def _aplatir_json(noeud, prefixe="", conteneur=None, index=None):
    """(chemin, valeur, conteneur, clé) — la clé permet de RÉÉCRIRE."""
    if isinstance(noeud, dict):
        for cle, valeur in noeud.items():
            yield from _aplatir_json(
                valeur, f"{prefixe}.{cle}".strip("."), noeud, cle
            )
    elif isinstance(noeud, list):
        for rang, valeur in enumerate(noeud):
            yield from _aplatir_json(valeur, f"{prefixe}[{rang}]", noeud, rang)
    else:
        yield prefixe, noeud, conteneur, index


def _lire_xml(chemin):
    """(feuilles, arbre). L'analyse désamorce les entités externes."""
    from defusedxml.ElementTree import parse

    arbre = parse(chemin)
    racine = arbre.getroot()
    lignes = [["chemin", "valeur"]]
    ancres = {}
    for element in racine.iter():
        if element.text and element.text.strip():
            lignes.append([element.tag, coercer_texte(element.text.strip())])
            ancres[(len(lignes), 2)] = (element, None)
        for cle, valeur in element.attrib.items():
            lignes.append([f"{element.tag}@{cle}", coercer_texte(valeur)])
            ancres[(len(lignes), 2)] = (element, cle)
        # La QUEUE d'un élément : le texte qui suit sa balise fermante. Un
        # export d'ERP nommé « .xls » qui est en réalité du HTML arrive
        # ici, et la moitié d'une cellule y vit — « Client <b>X</b> Nom »
        # porte « Nom » en queue de <b>.
        if element.tail and element.tail.strip():
            lignes.append(
                [f"{element.tag}#tail", coercer_texte(element.tail.strip())]
            )
            ancres[(len(lignes), 2)] = (element, "#tail")
    feuille = Feuille(NOM_FEUILLE_NEUTRE, lignes)
    feuille.ancres = ancres
    feuille.arbre = arbre
    feuille.colonne_structure = 1
    feuille.ligne1_fabriquee = True
    return [feuille], arbre


# ----------------------------------------------------------------------
# Le nettoyage hors cellules
# ----------------------------------------------------------------------
def nettoyer_hors_cellules(classeur, table, options):
    """Effacer ce qui porte du texte sans être une cellule. AVANT la grille.

    Sur les 24 vecteurs qu'un `.xlsx` peut porter, ce bloc en efface 19.
    Les cinq qui restent — plages nommées globale et locale, nom de
    tableau, littéral de formule, nom de feuille — sont référencés par des
    formules : les supprimer casserait ce que la règle de la formule vient
    de préserver, donc ils sont RAPPORTÉS et non effacés.

    La liste qui suit tient à un attribut privé près, et c'est le test de
    fuite de `test_transform_external.py` qui la garde : il assert la liste
    EXACTE des survivants, de sorte qu'un vecteur rouvert par une montée de
    version d'openpyxl fait tomber le test au lieu de passer inaperçu.
    """
    from openpyxl.packaging.core import DocumentProperties
    from openpyxl.packaging.custom import CustomPropertyList

    vivier = options["vivier"]
    comptes = {"proprietes": 0, "commentaires": 0, "hyperliens": 0}

    # `creator` vaut « openpyxl » par DÉFAUT : sans creator=None, l'élément
    # <dc:creator> ne disparaît pas, il est REMPLI. Mesuré.
    classeur.properties = DocumentProperties(creator=None)
    classeur.custom_doc_props = CustomPropertyList()
    comptes["proprietes"] += 1

    for onglet in classeur.worksheets:
        onglet._pivots = []
        onglet._images = []
        onglet.auto_filter.filterColumn = []
        onglet.data_validations.dataValidation = []
        onglet.conditional_formatting._cf_rules.clear()
        for entete in (
            onglet.oddHeader,
            onglet.evenHeader,
            onglet.firstHeader,
            onglet.oddFooter,
            onglet.evenFooter,
            onglet.firstFooter,
        ):
            for partie in (entete.left, entete.center, entete.right):
                partie.text = None
        if options.get("garder_graphiques"):
            _nettoyer_graphiques(onglet)
        else:
            onglet._charts = []
        for ligne in onglet.iter_rows():
            for cellule in ligne:
                if cellule.comment is not None:
                    cellule.comment = None
                    comptes["commentaires"] += 1
                if cellule.hyperlink is not None:
                    cellule.hyperlink = None
                    comptes["hyperliens"] += 1
        _anonymiser_noms_locaux(onglet, table, vivier)

    _anonymiser_noms_locaux(classeur, table, vivier)
    comptes["formats"] = _anonymiser_formats_de_nombre(classeur, table, vivier)
    comptes["styles"] = _renommer_styles_nommes(classeur)
    comptes["styles_tableau"] = _renommer_styles_de_tableau(classeur)
    comptes["polices"] = _renommer_polices(classeur)
    comptes["theme"] = _assainir_theme(classeur)
    return comptes


def _assainir_theme(classeur):
    """Le thème, que le graveur recopie octet pour octet.

    openpyxl rend `xl/theme/theme1.xml` tel qu'il l'a LU quand
    `loaded_theme` est rempli, et son propre défaut sinon. Deux chaînes
    libres y vivent : le nom sous lequel le thème a été enregistré, et les
    polices majeure et mineure. Une police de marque y reste donc nommée
    après que le renommage l'a retirée des styles, et une de ces chaînes
    qui répète une valeur de la grille rend le classeur inécrivable au
    filet, sans qu'aucune règle ne puisse l'assainir.

    Le remède est de laisser openpyxl écrire SON thème : la copie perd la
    palette du client, ce qui est exactement ce qu'on veut d'une copie
    transmissible.
    """
    if not getattr(classeur, "loaded_theme", None):
        return 0
    classeur.loaded_theme = None
    return 1


# Les caractères qui « avalent » le suivant dans un format de nombre :
# l'échappement, la réservation de largeur, la répétition. Un libellé peut
# s'écrire ainsi, hors guillemets et un caractère à la fois.
_ECHAPPE_FORMAT = "\\_*"

# Une section de devise : `[$USD-409]`, `[$-1010409]`. Son texte est libre —
# un nom de client y tient — alors que `[Red]`, `[<100]` et `[h]` sont des
# mots-clés qu'un remplacement casserait.
_DEVISE_CROCHET = re.compile(r"^\[\$(.*?)(-[0-9A-Fa-f]+)?\]$")

# Une section de devise qu'Excel écrit ENTIÈREMENT de convention : un
# symbole ou un code, suivi d'un modificateur de locale. Le seul LCID
# hexadécimal ne les couvre pas — « [$-en-US] », « [$-x-sysdate] »,
# « [$€-x-euro2] », « [$R$-pt-BR] ». Aucune ne porte de donnée du
# client, et les remplacer détruit le symbole monétaire de la copie,
# ou la forme de ses dates. Reconnue en ENTIER, jamais par un suffixe
# élargi : « [$Cabinet-Lav] » y perdrait la moitié de son libellé.
_DEVISE_LOCALE = re.compile(
    r"^\[\$[^\]-]{0,4}-(?:[0-9A-Fa-f]+|x-[a-z0-9]+"
    r"|[A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})*)\]$"
)


# Excel accepte du texte NU dans un format de nombre — sans guillemets ni
# barre oblique — et un nom écrit là sortait intact. Le reconnaître demande
# de distinguer un mot d'un motif de date, et les deux sont faits de
# lettres. Un ALPHABET ne le fait pas : Excel traduit ses marques de
# position dans la langue du classeur, si bien que « jj/mm/aaaa » et
# « tt.mm.jjjj » sont des dates aussi légitimes que « dd/mm/yyyy », et
# énumérer les lettres de chaque locale se perd — pour ensuite détruire la
# date de la copie sur celle qu'on a oubliée.
#
# Ce qui sépare les deux est la RÉPÉTITION. Une marque de position vient
# par groupes d'une même lettre, seuls ou collés — « aaaammjj », « hhmm ».
# Un mot colle des lettres différentes et retombe sur des groupes d'UNE
# seule — « Nom », « aboulie ». D'où le test : deux lettres différentes
# côte à côte ET un groupe solitaire.
#
# L'erreur reste orientée. Prendre un mot pour un motif le laisse passer,
# et le filet de relecture refuse la copie ; prendre un motif pour un mot
# abîme le classeur en silence.
_LETTRES_FORMAT = re.compile(r"[^\W\d_]+")

# Ce qu'Excel écrit en lettres sans que ce soit un libellé, malgré des
# lettres différentes côte à côte.
_MOTS_CLES_FORMAT = frozenset(("general", "am", "pm", "a", "p"))


def _est_un_mot(lettres):
    """Vrai si cette suite de lettres est un mot et non une marque."""
    if lettres.lower() in _MOTS_CLES_FORMAT:
        return False
    plie = lettres.lower()
    if not any(a != b for a, b in zip(plie, plie[1:])):
        return False
    return any(
        len(trouve.group()) < 2 for trouve in re.finditer(r"(.)\1*", plie)
    )


# Ce qui peut joindre deux mots d'un même libellé sans être du motif.
# Tout le reste sépare : bloquer la jonction ne coûte qu'un remplacement
# de plus, la forcer emporte le motif qui vit entre les deux.
_JOINT_LIBELLE = re.compile("^[\\s'\u2019-]*$")


def _spans_du_libelle(plage):
    """Les étendues des libellés nus de cette suite, de gauche à droite.

    Deux mots séparés d'un simple espace font UN libellé — « Nom du
    client » — mais deux mots séparés d'un motif en font DEUX : la
    tranche qui les réunirait emporterait ce motif, et « #,##0 X;-#,##0
    X » perdait sa section négative en entier.
    """
    spans = []
    for trouve in _LETTRES_FORMAT.finditer(plage):
        if not _est_un_mot(trouve.group()):
            continue
        debut, fin = trouve.span()
        if spans and _JOINT_LIBELLE.match(plage[spans[-1][1] : debut]):
            spans[-1] = (spans[-1][0], fin)
        else:
            spans.append((debut, fin))
    return spans


# Les jetons qu'une section entre crochets peut porter. Contrairement aux
# marques de position d'une date, la grammaire d'Excel les ÉNUMÈRE : huit
# couleurs, cinquante-six numérotées, une condition, une durée écoulée, un
# jeu de chiffres, une ère. C'est une spécification, pas une devinette,
# d'où l'énumération ici là où l'alphabet a été rejeté ailleurs.
#
# Tout jeton de trois caractères ou moins passe de toute façon, le seuil du
# filet le laissant : la liste ne sert que ceux qui l'atteignent.
_CROCHETS_CONNUS = frozenset(
    (
        "black",
        "blue",
        "cyan",
        "green",
        "magenta",
        "red",
        "white",
        "yellow",
        "thai",
        "hijri",
        "buddhist",
        "gregorian",
    )
)

_CROCHET_NUMEROTE = re.compile(
    r"^(?:color\s?(?:[1-9]|[1-4][0-9]|5[0-6])"
    r"|(?:db|nat)num(?:1[0-9]|[1-9]))$"
)

# Une condition, qu'une comparaison ouvre : `[<100]`, `[>=0]`, `[<>1]`.
_CROCHET_CONDITION = re.compile(r"^[<>=]")


def _crochet_est_connu(interieur):
    """Vrai si la grammaire d'Excel explique cette section."""
    plie = interieur.strip().lower()
    if len(plie) < noyau.LONGUEUR_VERIFIABLE:
        return True
    if _CROCHET_CONDITION.match(plie):
        return True
    return plie in _CROCHETS_CONNUS or bool(_CROCHET_NUMEROTE.match(plie))


def _parcourir_format(fmt, mot_si_long):
    r"""Le format parcouru de gauche à droite, ses libellés assainis.

    Excel a QUATRE façons de porter du texte dans un format de nombre, et
    une expression sur les seuls guillemets n'en voyait qu'une :

    - `"..."` — la forme courante ;
    - `\N\o\m` — une suite d'échappements, un caractère à la fois, sans
      aucun guillemet : c'est ainsi qu'un libellé écrit à la main dans
      Excel arrive souvent ;
    - `[$Nom-409]` — la section de devise, dont le texte est libre ;
    - un guillemet ÉCHAPPÉ, qui n'ouvre pas de littéral et décalait
      l'appariement de tous les suivants.

    Les autres sections entre crochets sont recopiées telles quelles : y
    remplacer `Red` ou `h` casserait le format.
    """
    sortie = []
    plat = []

    def vider_plat():
        """Le texte nu accumulé : ses libellés partent, son motif reste."""
        if not plat:
            return
        brut = "".join(plat)
        plat.clear()
        curseur = 0
        for debut, fin in _spans_du_libelle(brut):
            mot = mot_si_long(brut[debut:fin])
            if mot is None:
                continue
            sortie.append(brut[curseur:debut])
            sortie.append('"%s"' % mot)
            curseur = fin
        sortie.append(brut[curseur:])

    index, taille = 0, len(fmt)
    while index < taille:
        caractere = fmt[index]
        if caractere in _ECHAPPE_FORMAT and index + 1 < taille:
            vider_plat()
            debut = index
            lettres = []
            while index + 1 < taille and fmt[index] in _ECHAPPE_FORMAT:
                lettres.append(fmt[index + 1])
                index += 2
            mot = mot_si_long("".join(lettres))
            sortie.append(fmt[debut:index] if mot is None else '"%s"' % mot)
            continue
        if caractere == "[":
            vider_plat()
            fin = fmt.find("]", index)
            if fin == -1:
                sortie.append(fmt[index:])
                break
            section = fmt[index : fin + 1]
            if _DEVISE_LOCALE.match(section):
                sortie.append(section)
                index = fin + 1
                continue
            devise = _DEVISE_CROCHET.match(section)
            if devise:
                mot = mot_si_long(devise.group(1))
                if mot is not None:
                    section = "[$%s%s]" % (mot, devise.group(2) or "")
            elif not _crochet_est_connu(section[1:-1]):
                # Une section que la grammaire n'explique pas porte du
                # texte libre : Excel ne l'écrit pas, un producteur tiers
                # si, et elle sortait intacte. Le remplacement reste NU —
                # un crochet ne porte pas de guillemets, et la section
                # était déjà hors grammaire avant qu'on y touche.
                mot = mot_si_long(section[1:-1])
                if mot is not None:
                    section = "[%s]" % mot
            sortie.append(section)
            index = fin + 1
            continue
        if caractere == '"':
            vider_plat()
            fin = fmt.find('"', index + 1)
            interieur = fmt[index + 1 :] if fin == -1 else fmt[index + 1 : fin]
            mot = mot_si_long(interieur)
            sortie.append('"%s' % (interieur if mot is None else mot))
            if fin == -1:
                break
            sortie.append('"')
            index = fin + 1
            continue
        plat.append(caractere)
        index += 1
    vider_plat()
    return "".join(sortie)


def _anonymiser_formats_de_nombre(classeur, table, vivier):
    """Le texte libre d'un format de nombre personnalisé.

    Excel laisse suffixer un nombre d'un libellé — `#,##0" Nom du client"` —
    et ce libellé vit dans `xl/styles.xml`, hors de toute cellule. Aucune
    règle ne le voyait, et le filet le refusait sans jamais l'assainir : un
    classeur portant un nom dans un format personnalisé était inécrivable.

    La liste est REMPLACÉE dans son ordre, jamais réécrite cellule par
    cellule : les cellules référencent un format par son INDEX, et le
    setter d'openpyxl AJOUTE une entrée plutôt que de modifier la sienne —
    l'ancien format, littéral compris, repartait alors dans le fichier.

    Un littéral court — une devise, une unité — reste : il ne porte aucune
    donnée du client, et le remplacer abîmerait le classeur sans rien
    protéger. Le seuil est celui du filet, pour que ce qu'on garde ici soit
    exactement ce qu'il ne refusera pas.
    """
    from openpyxl.utils.indexed_list import IndexedList

    from script.data.external_file import nouveau_mot

    touches = 0

    def mot_si_long(interieur):
        """Le libellé assaini, ou None si le seuil du filet le laisse."""
        nonlocal touches
        noyau_texte = interieur.strip()
        if len(noyau_texte) < noyau.LONGUEUR_VERIFIABLE:
            return None
        touches += 1
        tete = interieur[: len(interieur) - len(interieur.lstrip())]
        queue = interieur[len(interieur.rstrip()) :]
        return "%s%s%s" % (
            tete,
            nouveau_mot(noyau_texte, table, vivier),
            queue,
        )

    formats = list(getattr(classeur, "_number_formats", []) or [])
    if formats:
        classeur._number_formats = IndexedList(
            [_parcourir_format(f, mot_si_long) for f in formats]
        )
    # Un style DIFFÉRENTIEL porte son propre format : c'est celui d'une mise
    # en forme conditionnelle. Vider la liste casserait l'index qu'un style
    # de tableau personnalisé y référence, donc on l'assainit en place.
    for style in (
        getattr(
            getattr(classeur, "_differential_styles", None), "styles", None
        )
        or []
    ):
        fmt = getattr(style, "numFmt", None)
        if fmt is not None and getattr(fmt, "formatCode", None):
            fmt.formatCode = _parcourir_format(fmt.formatCode, mot_si_long)
    for onglet in classeur.worksheets:
        for graphique in getattr(onglet, "_charts", []) or []:
            _formats_de_graphique(graphique, mot_si_long)
    return touches


def _formats_de_graphique(graphique, mot_si_long):
    """Les formats de nombre qu'un graphique porte hors des cellules.

    Un axe et une étiquette de données portent leur PROPRE format, dont le
    libellé vit dans `xl/charts/chartN.xml` : ni dans la liste que le
    classeur tient, ni dans une cellule. Le filet ne le rattrape pas, car
    il ne refuse que ce qui a été ANNONCÉ remplacé — un libellé qui n'a
    jamais été lu d'une cellule n'est annoncé par personne, et sort intact.

    Deux formes cohabitent : un objet à `formatCode` sur un axe, une
    chaîne nue sur une étiquette. Le parcours descend le graphe que
    openpyxl sérialise, et s'amorce sur les axes : le graphique ne les
    référence que par leur identifiant, ils ne sont pas dans ses éléments.
    """
    vus = set()
    pile = [graphique]
    for nom in ("x_axis", "y_axis", "z_axis"):
        axe = getattr(graphique, nom, None)
        if axe is not None:
            pile.append(axe)
    while pile:
        porteur = pile.pop()
        if id(porteur) in vus:
            continue
        vus.add(id(porteur))
        fmt = getattr(porteur, "numFmt", None)
        if isinstance(fmt, str):
            porteur.numFmt = _parcourir_format(fmt, mot_si_long)
        elif getattr(fmt, "formatCode", None):
            fmt.formatCode = _parcourir_format(fmt.formatCode, mot_si_long)
        for element in getattr(porteur, "__elements__", ()) or ():
            valeur = getattr(porteur, element, None)
            if not isinstance(valeur, (list, tuple)):
                valeur = (valeur,)
            for candidat in valeur:
                if hasattr(candidat, "__elements__"):
                    pile.append(candidat)


# Ce qu'installent Office, LibreOffice et les systèmes, plus les
# métriques-compatibles libres. Une police absente de cette liste devient
# `police_<n>` : la copie perd son apparence, ce qui est le côté sur
# lequel pencher quand le nom peut être celui d'une fonte de marque.
POLICES_COURANTES = frozenset(
    (
        "Aptos",
        "Aptos Display",
        "Aptos Narrow",
        "Arial",
        "Arial Black",
        "Arial Narrow",
        "Bahnschrift",
        "Bookman Old Style",
        "Cabin",
        "Calibri",
        "Calibri Light",
        "Cambria",
        "Cambria Math",
        "Candara",
        "Carlito",
        "Caladea",
        "Century Gothic",
        "Comic Sans MS",
        "Consolas",
        "Constantia",
        "Corbel",
        "Courier",
        "Courier New",
        "DejaVu Sans",
        "DejaVu Sans Mono",
        "DejaVu Serif",
        "Ebrima",
        "Franklin Gothic Book",
        "Garamond",
        "Georgia",
        "Gill Sans MT",
        "Helvetica",
        "Helvetica Neue",
        "Impact",
        "Inconsolata",
        "Lato",
        "Liberation Mono",
        "Liberation Sans",
        "Liberation Sans Narrow",
        "Liberation Serif",
        "Lucida Console",
        "Lucida Sans Unicode",
        "MS Gothic",
        "MS PGothic",
        "MS Sans Serif",
        "MS Serif",
        "Malgun Gothic",
        "Menlo",
        "Meiryo",
        "Monaco",
        "Noto Sans",
        "Noto Serif",
        "Open Sans",
        "Palatino Linotype",
        "PT Sans",
        "Roboto",
        "SimSun",
        "Segoe UI",
        "Segoe UI Light",
        "Segoe UI Semibold",
        "Segoe UI Symbol",
        "Source Sans Pro",
        "Sylfaen",
        "Symbol",
        "Tahoma",
        "Times",
        "Times New Roman",
        "Trebuchet MS",
        "Ubuntu",
        "Ubuntu Mono",
        "Verdana",
        "Webdings",
        "Wingdings",
        "Wingdings 2",
        "Wingdings 3",
        "Yu Gothic",
    )
)


def _renommer_polices(classeur):
    """Le NOM d'une police part aussi dans les styles.

    Une police installée chez le client porte son nom, et rien ne la
    référence par autre chose que ce nom. Un repère POSITIONNEL suffit —
    la faire passer par la table brûlerait des mots du vivier sur ce qui
    n'est pas une donnée de la grille, et ferait tolérer ce mot par le
    filet là où il n'a rien à excuser.

    Aucune forme ne sépare « Century Gothic » d'une fonte de marque : les
    deux sont des noms propres. Le renommage porte donc sur TOUTES celles
    que la liste ne nomme pas, et l'erreur penche du côté du dégât
    cosmétique plutôt que du nom qui sort. `POLICES_COURANTES` n'est pas
    la propriété de sûreté, seulement le confort de la copie : ce qu'elle
    oublie perd son apparence, jamais sa donnée.
    """
    touches = 0
    rang = 0
    # Un style DIFFÉRENTIEL porte sa police EN LIGNE : elle vit dans
    # `xl/styles.xml` sans passer par la liste des polices du classeur, si
    # bien qu'une fonte de marque survivait au renommage en n'étant nommée
    # que par une mise en forme conditionnelle.
    for porteur in list(getattr(classeur, "_fonts", []) or []) + [
        getattr(style, "font", None)
        for style in (
            getattr(
                getattr(classeur, "_differential_styles", None),
                "styles",
                None,
            )
            or []
        )
    ]:
        if porteur is None:
            continue
        rang += 1
        nom = getattr(porteur, "name", None)
        if nom and nom not in POLICES_COURANTES:
            porteur.name = f"police_{rang}"
            touches += 1
    return touches


def _renommer_styles_de_tableau(classeur):
    """Le NOM d'un style de tableau personnalisé, à ses TROIS endroits.

    Il vit dans les styles du classeur — l'entrée du style et, quand il est
    le défaut, l'attribut qui le désigne — et dans chaque partie de tableau
    qui le référence. Aucune formule ne le résout, contrairement au nom
    d'affichage d'un tableau : il se renomme, il ne se rapporte pas. Les
    trois endroits changent dans la même passe, sinon le tableau perd sa
    mise en forme.

    La liste ne porte que ce que le DOCUMENT définit : les styles intégrés
    n'y figurent jamais, donc chaque entrée est un nom tapé par quelqu'un.
    """
    liste = getattr(classeur, "_table_styles", None)
    correspondance = {}
    for rang, style in enumerate(
        getattr(liste, "tableStyle", None) or (), start=1
    ):
        if style.name:
            correspondance[style.name] = style.name = f"tableStyle_{rang}"
    if not correspondance:
        return 0
    for attribut in ("defaultTableStyle", "defaultPivotStyle"):
        valeur = getattr(liste, attribut, None)
        if valeur in correspondance:
            setattr(liste, attribut, correspondance[valeur])
    for onglet in classeur.worksheets:
        for tableau in (getattr(onglet, "tables", {}) or {}).values():
            info = getattr(tableau, "tableStyleInfo", None)
            nom = getattr(info, "name", None)
            if nom in correspondance:
                info.name = correspondance[nom]
    return len(correspondance)


def _renommer_styles_nommes(classeur):
    """Le NOM d'un style nommé part aussi dans `xl/styles.xml`.

    Renommé en place : une cellule référence son style par l'index de la
    liste, que renommer l'objet ne déplace pas. « Normal » est le style par
    défaut d'Excel et n'est pas un nom donné par quelqu'un.
    """
    correspondance = {}
    for rang, style in enumerate(
        getattr(classeur, "_named_styles", []) or [], start=1
    ):
        if style.name != "Normal":
            correspondance[style.name] = style.name = f"style_{rang}"
    if not correspondance:
        return 0
    # Une COLONNE de tableau référence un style nommé par son NOM — sur le
    # tableau et sur chacune de ses colonnes, pour l'en-tête, les données
    # et la ligne de total. Sans cette passe, l'ancien nom — celui du
    # client — reste dans la partie de tableau, et la référence ne résout
    # plus.
    for onglet in classeur.worksheets:
        for tableau in (getattr(onglet, "tables", {}) or {}).values():
            porteurs = [tableau] + list(
                getattr(tableau, "tableColumns", None) or []
            )
            for porteur in porteurs:
                for attribut in (
                    "headerRowCellStyle",
                    "dataCellStyle",
                    "totalsRowCellStyle",
                    "headerRowDxfId",
                ):
                    valeur = getattr(porteur, attribut, None)
                    if valeur in correspondance:
                        setattr(porteur, attribut, correspondance[valeur])
    return len(correspondance)


def _nettoyer_graphiques(onglet):
    """Le chemin optionnel : garder le graphique, vider ce qu'il porte.

    `s.tx = None` est INDISPENSABLE et le nettoyage des caches ne le
    remplace pas : un titre de série a deux formes — `<tx><v>littéral</v>`
    quand il est tapé, `<tx><strRef><f>réf</f>` quand il vient des données
    — et ni l'une ni l'autre n'est un cache : vider `strCache` et
    `numCache` les laisse toutes deux en place.
    """
    for graphique in onglet._charts:
        graphique.title = None
        for axe in (
            getattr(graphique, "x_axis", None),
            getattr(graphique, "y_axis", None),
            getattr(graphique, "z_axis", None),
        ):
            if axe is not None:
                axe.title = None
        for serie in graphique.series:
            for source in (
                serie.cat,
                serie.val,
                getattr(serie, "xVal", None),
                getattr(serie, "yVal", None),
                getattr(serie, "bubbleSize", None),
            ):
                if source is None:
                    continue
                for reference in ("numRef", "strRef", "multiLvlStrRef"):
                    porteur = getattr(source, reference, None)
                    if porteur is None:
                        continue
                    for cache in (
                        "numCache",
                        "strCache",
                        "multiLvlStrCache",
                    ):
                        if hasattr(porteur, cache):
                            setattr(porteur, cache, None)
            serie.tx = None


def _anonymiser_noms_locaux(porteur, table, vivier):
    """La constante littérale d'une plage nommée passe par la table.

    Le NOM ne se supprime pas — une formule le référencerait dans le vide,
    alors que la règle de la formule vient de la préserver. Sa VALEUR, en
    revanche, se remplace quand c'est une constante littérale, et le nom
    reste résolvable. Mesuré comme survivant tant qu'on ne le fait pas.
    """
    from script.data.external_file import nouveau_mot

    noms = getattr(porteur, "defined_names", None)
    if not noms:
        return
    for nom in list(noms.keys()):
        defini = noms[nom]
        texte = getattr(defini, "attr_text", "") or ""
        if len(texte) > 1 and texte.startswith('"') and texte.endswith('"'):
            interieur = texte[1:-1]
            if interieur:
                defini.attr_text = '"%s"' % nouveau_mot(
                    interieur, table, vivier
                )


# ----------------------------------------------------------------------
# La marche à blanc et l'écriture
# ----------------------------------------------------------------------
def _preparer(chemin, options):
    """(format, feuilles, classeur, rapport, options complétées)."""
    format_lu = detect_format(chemin)
    if format_lu == "protege":
        raise ErreurMoteur("protege", chemin)
    if format_lu == "xlsb":
        raise ErreurMoteur("illisible_ici", chemin)
    if not format_lu:
        raise ErreurMoteur("format_inconnu", os.path.basename(chemin))

    rapport = report(chemin)
    options = dict(options)
    options["vivier"] = vivier_de_mots()
    options.setdefault("nombres", True)
    options.setdefault("texte", True)
    options.setdefault("entetes", False)
    options["colonnes_intactes"] = set(options.get("colonnes_intactes") or [])
    # Un dict {nom: [str]} traverse le sous-processus ; un set et un dict à
    # clés tuple sont refusés par `json.dumps`, ce qui est la raison pour
    # laquelle les clés dérivées de `_preparer` ne voyagent jamais.
    options["colonnes_intactes_par_feuille"] = {
        str(nom): {str(c).strip() for c in (liste or []) if str(c).strip()}
        for nom, liste in (
            options.get("colonnes_intactes_par_feuille") or {}
        ).items()
    }

    classeur = None
    if format_lu == "xlsx":
        classeur, feuilles = _lire_xlsx(
            chemin, garder_vba=options.get("garder_macros")
        )
    elif format_lu == "xls":
        feuilles = _lire_xls(chemin)
    elif format_lu == "access":
        feuilles = _lire_access(chemin)
    elif format_lu == "csv":
        feuille_csv, meta_csv = _lire_csv(chemin)
        feuilles = [feuille_csv]
        # Le délimiteur est détecté, imprimé à l'opérateur, puis il servait
        # à LIRE et pas à écrire : un fichier à point-virgule revenait en
        # virgule, et le tableur du destinataire le rendait en une colonne.
        options.setdefault("delimiteur", meta_csv["delimiteur"])
    elif format_lu == "json":
        feuilles = _lire_json(chemin)[0]
    else:
        feuilles = _lire_xml(chemin)[0]

    # La même mesure que dans `report`, et la CORRECTION de l'opérateur
    # par-dessus : elle voyage dans les options sérialisées et ne se
    # remesure jamais. Posée avant `etiquettes` et `bornes`, qui en
    # dérivent par la ligne de champs.
    # La table est relue ici, et de nouveau par l'appelant : deux
    # lectures d'un fichier qui ne change pas entre les deux, ce qui
    # évite de faire traverser un objet à `_preparer` pour un
    # dictionnaire de quelques lignes.
    mesurer_entetes(
        feuilles,
        format_lu,
        options.get("entetes_par_feuille") or {},
        Correspondance.charger(options.get("table_chemin")).entetes,
    )
    connues = {f.nom for f in feuilles}
    demandees = options.get("feuilles") or []
    inconnues = [n for n in demandees if n not in connues]
    if inconnues:
        raise ErreurMoteur("aucune_feuille", ", ".join(inconnues))
    options["feuilles"] = list(demandees) or None
    options["etiquettes"] = _etiquettes_par_colonne(feuilles)
    options["bornes"] = _bornes_par_colonne(rapport)
    options["formes"] = _formes_par_colonne(rapport)
    options["formes_relation"] = _formes_par_colonne(rapport, "forme_relation")
    options["selections"] = _formes_par_colonne(rapport, "selection")
    # Hors tableur, la colonne 1 de la grille porte des NOMS de balise ou des
    # chemins de clé : de la structure, que le graveur ne touche jamais. Les
    # compter comme remplacées désarmait le refus « rien à faire » et brûlait
    # le vivier sur des noms de champ.
    _verifier_conversion(format_lu, feuilles, options)
    options["colonnes_structure"] = {
        (f.nom, f.colonne_structure) for f in feuilles if f.colonne_structure
    }
    # Hors tableur, l'ancre est la SEULE voie d'écriture, et aucune ancre ne
    # porte la ligne 1 : ni l'étiquette que le lecteur a fabriquée, ni les
    # clés d'objet d'un tableau d'enregistrements. La compter en portée la
    # fait annoncer comme remplacée alors que la copie la garde en clair, et
    # le filet qui relit les octets refuse alors TOUTE copie. Une conversion
    # sort par un autre graveur, qui écrit bien cette ligne : elle y reste
    # en portée.
    ecrit_par_ancres = (
        format_lu in ("xml", "json")
        and (options.get("conversion") or format_lu) == format_lu
    )
    options["lignes_structure"] = {
        (f.nom, 1) for f in feuilles if f.ligne1_fabriquee or ecrit_par_ancres
    }
    # L'empan d'en-tête, en couples, comme `lignes_structure` : bâti ICI,
    # donc jamais sérialisé. Une feuille sans en-tête n'y met rien, et sa
    # ligne 1 entre en portée — c'est le correctif de la fuite.
    options["lignes_entete"] = {
        (f.nom, n) for f in feuilles for n in (f.lignes_entete or ())
    }
    return format_lu, feuilles, classeur, rapport, options


def _verifier_conversion(format_lu, feuilles, options):
    """Refuser une cible impossible AVANT l'aperçu.

    Le refus venait du graveur, donc après que l'opérateur avait lu un
    aperçu propre et consenti — et sous une clé qui accusait le format
    au lieu de nommer la contrainte.
    """
    cible = options.get("conversion") or ""
    if not cible or cible == format_lu:
        return
    retenues = [
        f
        for f in feuilles
        if not options.get("feuilles") or f.nom in options["feuilles"]
    ]
    if cible == "xml" and len(retenues) > 1:
        raise ErreurMoteur("conversion_impossible", f"{len(retenues)} → xml")


# Combien de cellules d'empan gardées en clair sont LISTÉES par feuille.
# Par feuille, non globalement : un plafond global cachait une feuille
# entière derrière les entrées d'une autre. Il borne la LISTE, jamais le
# compte — ce qui dépasse est annoncé, faute de quoi l'opérateur consent
# sur un extrait qu'il prend pour le tout.
PLAFOND_GARDEES_LISTEES = 12


def _parcourir(
    feuilles, options, table, rng, appliquer=None, gardees_out=None
):
    """La passe unique : compte, et écrit si `appliquer` est donné.

    Un seul parcours pour la marche à blanc et pour l'écriture : deux
    chemins divergeraient, et c'est la marche à blanc qui perdrait la
    confiance qu'on lui accorde.
    """
    bilan = {
        "remplacees": 0,
        "texte": 0,
        "nombre": 0,
        "hors_portee": 0,
        "intactes": {},
        "apercu": [],
        "entete_gardee": [],
        "entete_gardee_omises": {},
        "colonnes_en_clair": [],
    }
    # Par feuille : TOUTES les cellules d'empan gardées, listées ou non.
    # Le recompte par balayage de la liste était quadratique, et ne
    # pouvait de toute façon pas compter au-delà du plafond.
    gardees_par_feuille = collections.Counter()
    # Par colonne : ce qu'elle porte, et ce qui y a été remplacé. Une
    # colonne qui porte quelque chose et dont RIEN n'a été remplacé part
    # entière en clair, et le critère vaut quelle que soit la raison —
    # plancher, réponse, famille qui traverse par règle, ou aucune règle
    # qui l'atteigne. Les compter par famille sur tout le fichier ne
    # nommait pas la colonne : sur quarante colonnes, « 7 date(s)
    # laissées » ne dit pas laquelle sort.
    contenu = collections.Counter()
    remplacees = collections.Counter()
    for feuille in feuilles:
        for numero, ligne in enumerate(feuille.lignes, start=1):
            for index, valeur in enumerate(ligne, start=1):
                famille = classer(valeur)
                if famille == "vide":
                    continue
                # Une ligne d'en-tête n'est pas du contenu de colonne :
                # la garder en clair est la RÈGLE, et la compter ici
                # ferait dire d'une colonne entièrement remplacée qu'elle
                # sort en clair.
                if (feuille.nom, numero) not in (
                    options.get("lignes_entete") or ()
                ):
                    contenu[(feuille.nom, index)] += 1
                if not cellule_en_portee(feuille.nom, numero, index, options):
                    bilan["hors_portee"] += 1
                    if gardees_out is not None:
                        _noter_gardee(gardees_out, feuille.nom, valeur)
                    if (
                        numero in (feuille.lignes_entete or ())
                        and not options.get("entetes")
                        # TOUTE cellule gardée, non les seules textuelles :
                        # l'empan garde aussi les lignes de mise en page
                        # au-dessus de la ligne de champs, et un nombre
                        # laissé là n'était dit à personne.
                        and famille != "formule"
                        and (feuille.nom, numero)
                        not in (options.get("lignes_structure") or ())
                    ):
                        # L'empan est MESURÉ, et il porte plusieurs lignes
                        # sur un export mis en page. Un compteur ne dirait
                        # pas qu'un nom de client est dedans : chaque
                        # valeur se montre, jusqu'au plafond.
                        gardees_par_feuille[feuille.nom] += 1
                        if (
                            gardees_par_feuille[feuille.nom]
                            <= PLAFOND_GARDEES_LISTEES
                        ):
                            bilan["entete_gardee"].append(
                                {
                                    "feuille": feuille.nom,
                                    # La VRAIE coordonnée : le littéral
                                    # « L1 » mentait dès que l'en-tête
                                    # n'était pas en ligne 1, et
                                    # l'opérateur cherchait la valeur au
                                    # mauvais endroit.
                                    "cellule": f"L{numero}C{index}",
                                    "valeur": valeur_hors_tableur(valeur),
                                }
                            )
                    continue
                bornes = options["bornes"].get((feuille.nom, index))
                if isinstance(valeur, (dict, list)):
                    # Un conteneur imbriqué n'est pas une cellule : sans
                    # cette récursion, tout ce qu'il porte sort en clair.
                    # Le compte des FEUILLES réécrites décide : reconstruit
                    # à l'identique, le conteneur reste intact.
                    neuve, compte = _transformer_json(
                        valeur, options, table, rng
                    )
                    if not compte:
                        neuve = _INTACTE
                else:
                    neuve = anonymise_cellule(
                        valeur, options, table, rng, bornes=bornes
                    )
                    compte = {famille: 1}
                if neuve is _INTACTE:
                    bilan["intactes"][famille] = (
                        bilan["intactes"].get(famille, 0) + 1
                    )
                    if gardees_out is not None:
                        _noter_gardee(gardees_out, feuille.nom, valeur)
                    continue
                remplacees[(feuille.nom, index)] += 1
                for fam, feuilles_reecrites in compte.items():
                    bilan["remplacees"] += feuilles_reecrites
                    if fam in ("texte", "nombre"):
                        bilan[fam] += feuilles_reecrites
                if len(bilan["apercu"]) < 5:
                    bilan["apercu"].append(
                        {
                            "feuille": feuille.nom,
                            "cellule": f"L{numero}C{index}",
                            "avant": valeur_hors_tableur(valeur),
                            "apres": valeur_hors_tableur(neuve),
                        }
                    )
                if appliquer is not None:
                    appliquer(feuille, numero, index, neuve)
    bilan["entete_gardee_omises"] = {
        nom: compte - PLAFOND_GARDEES_LISTEES
        for nom, compte in gardees_par_feuille.items()
        if compte > PLAFOND_GARDEES_LISTEES
    }
    bilan["colonnes_en_clair"] = [
        {
            "feuille": nom,
            "index": index,
            "etiquette": options["etiquettes"].get((nom, index)) or "",
            "cellules": porte,
        }
        for (nom, index), porte in sorted(contenu.items())
        if porte and not remplacees[(nom, index)]
    ]
    return bilan


def _nommer_les_colonnes_ecartees(bilan, rapport, options):
    """Les deux listes de colonnes, posées ensemble sur le bilan.

    `colonnes_ecartees` dit celles qu'une RÈGLE explique — le plancher, ou
    une réponse de l'opérateur. `colonnes_en_clair` garde alors les
    autres : celles qui sortent entières sans que rien de ce qui a été
    demandé ne le dise. Les laisser dans les deux listes les faisait
    annoncer deux fois, et une ligne redondante apprend à ne plus lire les
    autres.
    """
    bilan["colonnes_ecartees"] = _colonnes_ecartees(rapport, options)
    expliquees = {
        (c["feuille"], c["index"])
        for c in bilan["colonnes_ecartees"]
        if c.get("index")
    }
    bilan["colonnes_en_clair"] = [
        c
        for c in bilan.get("colonnes_en_clair") or []
        if (c["feuille"], c["index"]) not in expliquees
    ]


def _colonnes_ecartees(rapport, options):
    ecartees = []
    for feuille in rapport.get("feuilles", []):
        for colonne in feuille.get("colonnes", []):
            if colonne.get("plancher"):
                ecartees.append(
                    {
                        "feuille": feuille["nom"],
                        # L'INDEX autant que l'étiquette : c'est par lui
                        # qu'une colonne se recoupe avec celles qui
                        # sortent entières en clair, et sans lui aucune
                        # n'était reconnue comme déjà expliquée.
                        "index": colonne["index"],
                        "etiquette": colonne["etiquette"],
                        "raison": "plancher",
                    }
                )
            elif noyau.colonne_repondue(
                feuille["nom"],
                colonne["index"],
                colonne.get("etiquette"),
                options,
            ):
                # La MÊME règle que la portée : la répéter ici faisait
                # trois occasions de divergence. Une colonne sans
                # étiquette ne se désigne que par son index, et l'aperçu
                # ne la nommait pas du tout.
                ecartees.append(
                    {
                        "feuille": feuille["nom"],
                        "index": colonne["index"],
                        "etiquette": (
                            colonne.get("etiquette") or f"#{colonne['index']}"
                        ),
                        "raison": "question",
                    }
                )
    return ecartees


def plan(chemin, options):
    """Ce qui SERAIT écrit. Rien n'est ouvert en écriture.

    La convention du dépôt pour l'action de menu qui écrit : montrer, puis
    demander. Une question posée avant de savoir ce qui sera touché n'est
    pas un consentement.
    """
    import random

    format_lu, feuilles, _classeur, rapport, options = _preparer(
        chemin, options
    )
    table = Correspondance.charger(options.get("table_chemin"))
    graine = options.get("graine")
    rng = (
        random.Random(graine) if graine not in (None, "") else random.Random()
    )
    fichiers = _fichiers_prevus(chemin, feuilles, options, table)
    bilan = _parcourir(feuilles, options, table, rng)
    _nommer_les_colonnes_ecartees(bilan, rapport, options)
    bilan["format"] = format_lu
    bilan["fichiers"] = fichiers
    # La marche à blanc est ce sur quoi l'opérateur consent : elle doit
    # refuser là où l'écriture refuserait.
    _refuser_table_confondue(
        chemin,
        options.get("table_chemin"),
        bilan["fichiers"] + [options.get("destination")],
    )
    bilan["avertissements"] = _avertissements(rapport, options)
    return bilan


def _fichiers_prevus(chemin, feuilles, options, table=None):
    """Les chemins qui seront écrits — calculés AVANT toute écriture.

    Les noms passent par la MÊME table que la conversion : un fichier par
    feuille est nommé d'après le nom ANONYMISÉ de la feuille, et prédire
    d'après le nom d'origine annonçait des chemins qui n'existeraient
    jamais. Ce n'était pas qu'un affichage — cette liste est ce que le
    contrôle d'écrasement et celui de la table de correspondance
    examinent, si bien qu'un fichier réel échappait aux deux.

    L'appel doit précéder le parcours des cellules, dans la marche à
    blanc comme à l'écriture : la table sert les deux, et le nom réservé
    ici est celui que la conversion retrouvera.
    """
    destination = options.get("destination") or ""
    cible = options.get("conversion") or ""
    if cible and _un_fichier_par_feuille(cible, feuilles, options):
        pris = set()
        return [
            os.path.join(
                destination,
                "%s.%s"
                % (
                    nom_de_fichier_sur(
                        _nom_de_feuille_anonyme(f, table, options), pris
                    ),
                    cible,
                ),
            )
            for f in feuilles
            if not options["feuilles"] or f.nom in options["feuilles"]
        ]
    return [destination] if destination else []


def _un_fichier_par_feuille(cible, feuilles, options):
    retenues = [
        f
        for f in feuilles
        if not options["feuilles"] or f.nom in options["feuilles"]
    ]
    return cible in ("csv",) and len(retenues) > 1


# Le plafond de `distinctes` dans `_stats_colonnes`. Au-delà, le compte
# ne dit plus combien la colonne porte de valeurs, et aucune conclusion
# sur sa saturation ne tient.
PLAFOND_DISTINCTES = 10000

# À partir de quelle part de l'étendue une colonne d'entiers est dite
# saturée. Pleine, elle ne laisse AUCUNE liberté au tirage ; à neuf
# dixièmes, la copie porte déjà presque le même ensemble.
PART_SATUREE = 0.9

# En deçà, une colonne d'entiers pleine ne dit rien de personne : un
# drapeau à deux états, un mois sur douze. L'avertissement y serait du
# bruit, et le bruit finit par se lire comme du fond.
SATURATION_MINIMALE = 20


def _colonnes_saturees(rapport, options):
    """Les colonnes d'entiers dont la copie sera une PERMUTATION.

    Le tirage est sans remise et reste dans l'étendue mesurée de la
    colonne : avec autant de valeurs distinctes que l'étendue compte
    d'entiers, l'ensemble de sortie est forcément l'ensemble d'entrée, et
    seule l'affectation change. Ce n'est pas une fuite — la permutation ne
    s'inverse pas sans la table — mais l'écran annonce « N nombres
    remplacés » et une comparaison d'ENSEMBLES ne montrerait rien.

    Une colonne planchéiée ou laissée intacte n'entre pas : elle n'est pas
    remplacée du tout, et `colonnes_ecartees` la nomme déjà.
    """
    if not options.get("nombres", True):
        return []
    saturees = []
    for feuille in rapport.get("feuilles", []):
        if (
            options.get("feuilles")
            and feuille["nom"] not in options["feuilles"]
        ):
            continue
        for colonne in feuille.get("colonnes", []):
            if colonne.get("plancher") or not colonne.get("entiere"):
                continue
            etiquette = colonne.get("etiquette")
            if noyau.colonne_repondue(
                feuille["nom"], colonne["index"], etiquette, options
            ):
                continue
            distinctes = colonne.get("distinctes") or 0
            if distinctes >= PLAFOND_DISTINCTES:
                continue
            if distinctes < SATURATION_MINIMALE:
                continue
            etendue = int(colonne["max"]) - int(colonne["min"]) + 1
            if etendue > 0 and distinctes >= etendue * PART_SATUREE:
                saturees.append((feuille["nom"], etiquette))
    return saturees


def _noms_de_feuille_survivent(rapport, options):
    """Le graveur recopie-t-il les noms d'onglet TELS QUELS ?

    Un seul chemin le fait : une source `.xlsx` rendue en `.xlsx` par
    `classeur.save`. Toute conversion passe par un classeur NEUF dont les
    onglets reçoivent un nom de la table, et `.xls` comme Access n'ont pas
    de graveur — leur copie repart par cette même conversion.

    Le dire quand ce n'est pas vrai n'est pas anodin : la liste des
    avertissements EST la surface du consentement, et un avis qui parle
    d'un risque écarté apprend à ne plus la lire.
    """
    if rapport.get("format") != "xlsx":
        return False
    cible = options.get("conversion") or ""
    return not (cible and cible != "xlsx")


def _avertissements(rapport, options):
    """Ce que la copie perd ou garde, dit plutôt que découvert."""
    dits = []
    if _colonnes_saturees(rapport, options):
        dits.append(
            "An integer column is saturated: the copy holds the same set"
            " of values, only reshuffled."
        )
    hors = rapport.get("hors_cellules") or {}
    if rapport.get("format") == "xlsx":
        dits.append(
            "Cached formula results are dropped;"
            " the sheet recomputes on open."
        )
        dits.append("Document properties were cleared on the copy.")
    if hors.get("croises") or hors.get("graphiques"):
        dits.append(
            "Pivot tables and chart caches are removed: they hold an"
            " unanonymised copy of the source."
        )
    if hors.get("images"):
        dits.append("Images and drawings are not carried over to the copy.")
    if hors.get("graphiques") and not options.get("garder_graphiques"):
        dits.append("Charts are not carried over to the copy.")
    if hors.get("graphiques") and options.get("garder_graphiques"):
        dits.append(
            "Charts are kept: their caches are cleaned through private"
            " attributes."
        )
    if hors.get("liens_externes"):
        dits.append(
            "External links were dropped; formulas that used them"
            " show #REF!."
        )
    survivent = _noms_de_feuille_survivent(rapport, options)
    if hors.get("plages_nommees") and survivent:
        dits.append(
            "Range and table names are kept so formulas resolve;"
            " they may hold identifying strings."
        )
    if survivent:
        # Le nom d'onglet survit dans workbook.xml, qu'une formule le
        # référence ou non : conditionner cet avertissement à la présence
        # d'un littéral de formule le taisait sur le cas le plus courant.
        dits.append(
            "Sheet names are kept so formulas resolve; they may identify."
        )
    if options.get("feuilles"):
        dits.append(
            "Sheets outside the selection are dropped from the copy;"
            " formulas that referenced them show #REF!."
        )
    if hors.get("feuilles_graphiques"):
        dits.append(
            "Chart sheets are dropped from the copy: their titles and"
            " series caches are not cells."
        )
    if rapport.get("format") == "json":
        dits.append("Object keys are kept as structure; they may identify.")
    if rapport.get("format") == "xml":
        dits.append(
            "Element and attribute names are kept as structure;"
            " they may identify."
        )
    encodage = (rapport.get("encodage") or "").lower().replace("-", "_")
    if encodage and encodage not in ("utf_8", "utf8", "ascii"):
        # Le délimiteur de la source est repris, son encodage NON : la
        # copie sort en UTF-8. C'est le bon choix — un mot du vivier ou un
        # en-tête gardé peut ne pas s'encoder dans le jeu d'origine, et
        # l'écriture échouerait après la question du consentement. Le
        # défaut était de ne pas le dire, alors que le rapport annonce
        # l'encodage détecté et laisse croire qu'il est conservé.
        dits.append("The copy is written in UTF-8, whatever the source was.")
    if options.get("garder_macros"):
        dits.append(
            "The VBA project and its companions (form controls, ActiveX,"
            " VML shapes, ribbon, EMF images) are copied as they are and"
            " were not reviewed."
        )
    return dits


def _ecrire_atomique(destination, ecrivain):
    """Un temporaire du même répertoire, puis `os.replace()`.

    Le `try/finally` n'est pas une politesse : sans lui, une exception
    laisse de la donnée client sur le disque, sous un nom que personne
    n'annonce. Le mode 0600 dès la création, parce que le temporaire porte
    cette donnée avant d'être renommé.
    """
    parent = os.path.dirname(os.path.abspath(destination)) or "."
    os.makedirs(parent, mode=0o700, exist_ok=True)
    descripteur, temporaire = tempfile.mkstemp(
        dir=parent, prefix=".transform-", suffix=".part"
    )
    os.close(descripteur)
    os.chmod(temporaire, 0o600)
    try:
        ecrivain(temporaire)
        os.replace(temporaire, destination)
        temporaire = None
    finally:
        if temporaire and os.path.exists(temporaire):
            os.unlink(temporaire)


def _refuser_si_source(chemin, destination):
    """La destination ne peut JAMAIS être la source.

    C'est le seul chemin par lequel l'original disparaîtrait, et la
    confirmation par nom — qui porte sur « le fichier existe déjà » — ne le
    distinguerait pas d'un écrasement ordinaire. `samefile` couvre le lien
    symbolique, `realpath` le couvre avant qu'il existe.
    """
    if os.path.exists(destination):
        try:
            if os.path.samefile(chemin, destination):
                raise ErreurMoteur("destination_source", destination)
        except OSError:
            pass
    if os.path.realpath(chemin) == os.path.realpath(destination):
        raise ErreurMoteur("destination_source", destination)


# Ce qu'une paire de crochets enferme dans une formule, sans les paires
# imbriquées : « [[#This Row],[Montant]] » rend « #This Row » et
# « Montant ».
_CROCHETS_FORMULE = re.compile(r"\[([^\[\]]*)\]")


def _litteraux_de_formule(formule, sans_egal=False):
    """Ce qu'une formule PRÉSERVE, et que le filet doit donc excuser.

    Les guillemets ne suffisent pas : une référence structurée de tableau
    porte le nom de colonne entre CROCHETS et sans guillemets —
    `=T[[#This Row],[Montant]]`. Toutes les cellules d'une colonne
    calculée partagent le MÊME texte, si bien que le bloc toléré ne
    l'excuse qu'une fois là où la copie le porte deux fois : dans la
    feuille et dans `xl/tables/`. Rendu en tolérance EXACTE, le nom cesse
    de se compter.

    Une formule matricielle n'est pas une `str` mais un objet à `text` ;
    la tester par `startswith` la faisait passer inaperçue.

    `sans_egal` pour un texte que l'appelant SAIT être une formule : OOXML
    omet le « = » initial dans `calculatedColumnFormula`, et la garde le
    rejetait donc — un tableau dont la colonne calculée n'a pas de formule
    de cellule équivalente se faisait refuser.
    """
    texte = (
        formule if isinstance(formule, str) else getattr(formule, "text", None)
    )
    # La garde sur « = » est LOAD-BEARING : appelée sur toute cellule, la
    # fonction tolérerait chaque valeur texte et le filet ne refuserait
    # plus rien.
    if not isinstance(texte, str) or not texte:
        return ()
    if not sans_egal and not texte.startswith("="):
        return ()
    morceaux = {texte}
    morceaux.update(re.findall(r'"([^"]*)"', texte))
    for brut in _CROCHETS_FORMULE.findall(texte):
        morceaux.add(brut)
        # Excel échappe `[`, `]`, `#` et l'apostrophe d'un nom de colonne
        # par une apostrophe : le nom réel est la forme déséchappée.
        morceaux.add(brut.replace("'", ""))
    return {m for m in morceaux if m}


def _valeurs_gardees(classeur, feuilles, format_lu):
    """Ce que le moteur conserve SCIEMMENT, et qui a été annoncé.

    La vérification d'après écriture refuse tout ce qui subsiste ; cette
    liste est la seule tolérance, et elle doit rester courte et justifiée.
    Chaque entrée est conservée parce que la retirer casserait ce que la
    règle de la formule vient de préserver.
    """
    gardees = set()
    for feuille in feuilles or ():
        # `ligne[0]` n'est de la structure que là où il porte VRAIMENT un
        # chemin de clé. La colonne 1 d'un tableau d'objets porte des
        # VALEURS, et les tolérer aveuglait le filet sur la colonne même
        # qu'il doit surveiller.
        if feuille.colonne_structure != 1:
            continue
        for ligne in feuille.lignes:
            if ligne and isinstance(ligne[0], str):
                gardees.add(ligne[0])
                gardees.update(re.split(r"[.\[\]@#]", ligne[0]))
    if classeur is None:
        return gardees
    for onglet in classeur.worksheets:
        gardees.add(onglet.title)
        gardees.update(onglet.defined_names.keys())
        for tableau in getattr(onglet, "tables", {}) or {}:
            gardees.add(str(tableau))
        # Une colonne de tableau porte sa PROPRE formule, que rien ne lit
        # par la grille : elle vit dans `xl/tables/` et la règle de la
        # formule la préserve, donc ce qu'elle nomme doit être excusé.
        for tableau in (getattr(onglet, "tables", {}) or {}).values():
            for tcol in getattr(tableau, "tableColumns", None) or ():
                for attribut in (
                    "calculatedColumnFormula",
                    "totalsRowFormula",
                ):
                    gardees.update(
                        _litteraux_de_formule(
                            getattr(tcol, attribut, None), sans_egal=True
                        )
                    )
        for ligne in onglet.iter_rows():
            for cellule in ligne:
                # Le TEXTE d'une formule : la règle 1 interdit d'y toucher,
                # et le rapport le compte et l'annonce.
                gardees.update(_litteraux_de_formule(cellule.value))
    gardees.update(classeur.defined_names.keys())
    return gardees


def ecrire(chemin, destination, options):
    """Écrire la copie. L'original n'est jamais modifié."""
    import random

    _refuser_si_source(chemin, destination)
    format_lu, feuilles, classeur, rapport, options = _preparer(
        chemin, options
    )
    options["destination"] = destination
    table = Correspondance.charger(options.get("table_chemin"))
    prevus = _fichiers_prevus(chemin, feuilles, options, table)
    _refuser_table_confondue(
        chemin, options.get("table_chemin"), prevus + [destination]
    )
    graine = options.get("graine")
    rng = (
        random.Random(graine) if graine not in (None, "") else random.Random()
    )
    hors_cellules = {}
    if format_lu == "xlsx":
        hors_cellules = nettoyer_hors_cellules(classeur, table, options)

    def appliquer(feuille, numero, index, neuve):
        feuille.lignes[numero - 1][index - 1] = neuve
        if feuille.source is not None:
            feuille.source.cell(row=numero, column=index).value = neuve
        # Hors tableur, l'ancre est la SEULE voie d'écriture : sans elle,
        # le graveur reparcourait l'arbre et ignorerait la portée.
        ancre = feuille.ancres.get((numero, index))
        if ancre is None:
            return
        conteneur, cle = ancre
        if isinstance(conteneur, (dict, list)):
            conteneur[cle] = neuve
        elif cle is None:
            conteneur.text = _garder_espaces(conteneur.text, neuve)
        elif cle == "#tail":
            conteneur.tail = _garder_espaces(conteneur.tail, neuve)
        else:
            conteneur.attrib[cle] = "" if neuve is None else str(neuve)

    hors_portee_gardees = {}
    bilan = _parcourir(
        feuilles,
        options,
        table,
        rng,
        appliquer=appliquer,
        gardees_out=hors_portee_gardees,
    )
    if not bilan["remplacees"] and not options.get("conversion"):
        if options["colonnes_intactes"] or not (
            options["nombres"] or options["texte"]
        ):
            raise ErreurMoteur("tout_exclu", "")
        raise ErreurMoteur("rien_a_faire", "")

    cible = options.get("conversion") or ""
    if format_lu == "xlsx":
        _resynchroniser_tableaux(classeur)
        # Récolter les tolérances sur le classeur qui SERA enregistré : un
        # littéral de formule ou un titre pris sur une feuille que la copie
        # ne porte pas taisait le filet au nom d'une formule que le
        # destinataire ne verra jamais.
        if not (cible and cible != format_lu):
            _retirer_feuilles_hors_portee(classeur, options)
    gardees = _valeurs_gardees(classeur, feuilles, format_lu)
    # Une cellule hors portée reste EN CLAIR par décision, et l'aperçu la
    # nomme : sa valeur n'est pas une fuite. Seules comptent les feuilles
    # qui partent dans la copie — celles que `_retirer_feuilles_hors_portee`
    # supprime n'y sont plus, et une de leurs valeurs qui reparaîtrait dans
    # un cache est une fuite comme une autre.
    retenues = options.get("feuilles")
    for nom, valeurs_vues in hors_portee_gardees.items():
        if retenues and nom is not None and nom not in retenues:
            continue
        gardees.update(valeurs_vues)

    if cible and cible != format_lu:
        fichiers = convertir(
            chemin,
            destination,
            cible,
            options,
            feuilles=feuilles,
            table=table,
        )
    elif format_lu == "xlsx":
        _ecrire_atomique(destination, classeur.save)
        fichiers = [destination]
    elif format_lu == "csv":
        fichiers = _ecrire_csv(destination, feuilles[0], options)
    elif format_lu in ("json", "xml"):
        fichiers = _ecrire_arbre(destination, feuilles[0], format_lu)
    else:
        # `.xls` et Access n'ont pas de graveur : la copie repart en .xlsx
        fichiers = convertir(
            chemin,
            destination,
            "xlsx",
            options,
            feuilles=feuilles,
            table=table,
        )

    # Le filet : relire les OCTETS écrits. Il ne dépend d'aucune
    # énumération de vecteurs, donc un endroit du format que personne n'a
    # pensé à nettoyer produit un refus, là où une liste de parties à
    # vérifier produirait un silence.
    fuites, non_vues = verifier_copie(fichiers, table, gardees)
    if fuites:
        for fichier in fichiers:
            if fichier and os.path.isfile(fichier):
                os.unlink(fichier)
        apercu = "; ".join(
            f"{v!r} -> {', '.join(sorted(set(parties))[:2])}"
            for v, parties in sorted(fuites.items())[:5]
        )
        # Une macro cite couramment la valeur d'une cellule, et le projet
        # VBA est recopié tel quel : garder les macros rend alors la copie
        # irrecevable. Le refus est juste, mais l'option qui le cause vient
        # d'être choisie une écran plus tôt, et le détail seul nomme une
        # partie du format sans dire quoi en faire.
        conseil = ""
        if options.get("garder_macros") and any(
            "vba" in partie.lower() or partie.lower().endswith(".bin")
            for parties in fuites.values()
            for partie in parties
        ):
            conseil = (
                "The kept VBA project quotes a source value;"
                " answer no to the macro question to write the copy."
            )
        raise ErreurMoteur(
            "fuite_detectee", f"{len(fuites)} — {apercu}", conseil
        )

    chemin_table = options.get("table_chemin")
    if chemin_table:
        try:
            # Les feuilles que l'opérateur a VRAIMENT changées, et
            # celles-là seules. `entetes_par_feuille` porte AUSSI les
            # empans mesurés — c'est ce qui a été montré à l'écran, et
            # c'est là-dessus qu'il a consenti — mais les retenir
            # imposerait la mesure d'un fichier à tout le lot, jusque là
            # où la mesure du suivant dirait autre chose. Sans cette
            # liste, une erreur de mesure devenait la loi du lot.
            demandees = options.get("entetes_par_feuille") or {}
            for nom in options.get("entetes_corrigees") or ():
                if str(nom) not in demandees:
                    continue
                table.entetes[str(nom)] = noyau.lignes_entieres(
                    demandees[str(nom)]
                )
            table.ecrire(chemin_table)
        except OSError:
            # Une copie sans sa table reçoit les mêmes mots que le fichier
            # suivant du lot, et deux clients fusionnent sur un seul mot.
            for fichier in fichiers:
                if fichier and os.path.isfile(fichier):
                    os.unlink(fichier)
            raise

    bilan["fichiers"] = fichiers
    bilan["hors_cellules"] = hors_cellules
    _nommer_les_colonnes_ecartees(bilan, rapport, options)
    bilan["avertissements"] = _avertissements(rapport, options)
    bilan["table"] = chemin_table or ""
    bilan["valeurs_non_verifiees"] = non_vues
    return bilan


def _garder_espaces(brut, neuve):
    """Réécrire un texte sans perdre son encadrement.

    Le texte d'un élément XML porte l'indentation du document et l'espace
    qui sépare deux mots d'un contenu mixte. Écrire la valeur nue collait
    « acai<b> » et aplatissait le document.
    """
    if neuve is None:
        return ""
    brut = brut or ""
    tete = brut[: len(brut) - len(brut.lstrip())]
    queue = brut[len(brut.rstrip()) :]
    return f"{tete}{neuve}{queue}"


def _retirer_feuilles_hors_portee(classeur, options):
    """Une feuille écartée de la portée ne part pas dans la copie.

    `classeur.save` grave le classeur ENTIER : répondre « Ventes » à la
    question des feuilles laissait les autres — dont une feuille masquée —
    intactes dans le fichier livré, avec pour seule trace « N cellules
    laissées hors portée ». Les autres cibles les retirent déjà.
    """
    # Une feuille graphique n'a AUCUNE cellule : la passe sur la grille ne
    # la voit pas, `worksheets` l'exclut par construction, et son titre,
    # ses titres d'axes, son en-tête et ses caches de série ne passent donc
    # par aucune règle. Rien ne peut l'anonymiser — elle part, toujours, et
    # l'avertissement le dit. La retirer seulement quand une sélection de
    # feuilles existe faisait mentir cet avertissement dans tous les autres
    # cas.
    for onglet in list(getattr(classeur, "chartsheets", []) or []):
        classeur.remove(onglet)

    retenues = options.get("feuilles")
    if not retenues:
        return
    for nom in [
        onglet.title
        for onglet in classeur.worksheets
        if onglet.title not in retenues
    ]:
        del classeur[nom]


def _noter_gardee(gardees_out, feuille, valeur):
    """Une valeur laissée intacte, notée sous ses DEUX formes.

    Le filet cherche la valeur telle que la table la porte, espaces
    compris ; ne noter que la forme dépouillée laissait un libellé à
    espace final annoncé remplacé ailleurs et jamais excusé ici — un
    refus sur du travail légitime, sur un fichier ordinaire dont un
    en-tête finit par une espace.
    """
    texte = str(valeur_hors_tableur(valeur) or "")
    for forme in (texte, texte.strip()):
        if len(forme) >= noyau.LONGUEUR_VERIFIABLE:
            gardees_out.setdefault(feuille, set()).add(forme)


def _resynchroniser_tableaux(classeur):
    """Un tableau porte une COPIE du texte de sa cellule d'en-tête.

    OOXML exige que `tableColumn.name` égale la cellule d'en-tête. Après
    l'anonymisation de celle-ci les deux divergent : rien n'est plus
    préservé au bénéfice d'une formule, et le texte d'origine reste
    simplement derrière. `totalsRowLabel` n'est jamais une cellule et
    survivait quelles que soient les options.

    Appelé APRÈS la passe sur la grille, contrairement au nettoyage hors
    cellules : il lui faut la valeur d'en-tête déjà remplacée.
    """
    for onglet in classeur.worksheets:
        for tableau in list((getattr(onglet, "tables", {}) or {}).values()):
            debut = tableau.ref.split(":")[0]
            colonne = onglet[debut].column
            # La DERNIÈRE ligne d'en-tête porte les noms de champ. Prendre
            # la première donnait, sur un en-tête de deux lignes, la ligne
            # de CATÉGORIE — souvent la même valeur sur plusieurs colonnes,
            # d'où des `tableColumn` de nom identique, qu'OOXML interdit :
            # Excel annonce alors un fichier à réparer.
            hauteur = getattr(tableau, "headerRowCount", 1)
            ligne = onglet[debut].row + max(int(hauteur or 1), 1) - 1
            pris = set()
            for rang, tcol in enumerate(tableau.tableColumns):
                cellule = onglet.cell(row=ligne, column=colonne + rang)
                nom = cellule.value
                nom = nom if isinstance(nom, str) and nom.strip() else ""
                # Un nom qu'AUCUNE cellule ne porte — en-tête vide, ou
                # tableau déclaré sans ligne d'en-tête — restait tel quel :
                # le texte d'origine sortait dans `xl/tables/`, et le filet
                # ne pouvait pas le voir puisqu'il n'a jamais été lu d'une
                # cellule, donc jamais annoncé remplacé.
                if not nom or nom in pris:
                    nom = "colonne_%d" % (rang + 1)
                pris.add(nom)
                tcol.name = nom
                if getattr(tcol, "totalsRowLabel", None):
                    tcol.totalsRowLabel = None


def _refuser_table_confondue(chemin, chemin_table, fichiers):
    """La table ne peut être NI la source NI un fichier écrit.

    Elle est gravée APRÈS la copie, en écrasant sa cible : confondue avec
    la source, elle détruit l'original ; confondue avec la copie, elle
    laisse sous le nom de la copie la liste en clair de toutes les valeurs
    réelles, et le menu annonce « Written: <destination> ».
    """
    if not chemin_table:
        return
    reel = os.path.realpath(chemin_table)
    if reel == os.path.realpath(chemin):
        raise ErreurMoteur("table_source", chemin_table)
    for fichier in fichiers:
        if fichier and reel == os.path.realpath(fichier):
            raise ErreurMoteur("table_source", chemin_table)


def _ecrire_arbre(destination, feuille, format_lu):
    """Graver l'arbre JSON ou XML que les ancres viennent de modifier.

    Aucun second parcours : ce que `_parcourir` a décidé est déjà DANS
    l'arbre, par les ancres. Un graveur qui reparcourait la source
    ignorerait la portée, le plancher et les colonnes intactes.
    """
    arbre = feuille.arbre

    def ecrivain(cible):
        if format_lu == "json":
            # Le porteur d'une racine scalaire : l'ancre a écrit DANS sa
            # case, c'est elle qu'il faut sérialiser. Pour un dict ou une
            # liste il vaut None, et `arbre` est déjà l'objet muté.
            charge = arbre if feuille.porteur is None else feuille.porteur[0]
            with open(cible, "w", encoding="utf-8") as fh:
                json.dump(charge, fh, ensure_ascii=False, indent=1)
        else:
            arbre.write(cible, encoding="utf-8", xml_declaration=True)

    _ecrire_atomique(destination, ecrivain)
    return [destination]


def _ecrire_csv(destination, feuille, options):
    def ecrivain(cible):
        with open(cible, "w", encoding="utf-8", newline="") as fh:
            graveur = csv.writer(
                fh, delimiter=options.get("delimiteur") or ","
            )
            for ligne in feuille.lignes:
                graveur.writerow(
                    [
                        "" if v is None else valeur_hors_tableur(v)
                        for v in ligne
                    ]
                )

    _ecrire_atomique(destination, ecrivain)
    return [destination]


def _transformer_json(noeud, options, table, rng):
    """(valeur, compte). Les CLÉS sont de la structure et restent.

    `compte` porte, par famille, le nombre de FEUILLES réécrites. Il est
    vide quand le sous-arbre revient identique : l'appelant en fait alors un
    « intact ». Sans ce compte, un conteneur reconstruit à l'identique
    valait une cellule remplacée — ce qui désarmait les refus « rien à
    faire » et « tout exclu », et faisait écrire une copie identique à la
    source en l'annonçant anonymisée.
    """
    if isinstance(noeud, (dict, list)):
        compte = {}
        entrees = (
            noeud.items() if isinstance(noeud, dict) else enumerate(noeud)
        )
        neuf = {} if isinstance(noeud, dict) else [None] * len(noeud)
        for cle, valeur in entrees:
            neuf[cle], partiel = _transformer_json(valeur, options, table, rng)
            for fam, feuilles_reecrites in partiel.items():
                compte[fam] = compte.get(fam, 0) + feuilles_reecrites
        return neuf, compte
    neuve = anonymise_cellule(noeud, options, table, rng)
    if neuve is _INTACTE:
        return noeud, {}
    return neuve, {classer(noeud): 1}


def _nom_de_feuille_anonyme(feuille, table, options):
    """Le nom d'une feuille, remplacé, pour une CONVERSION.

    Une table Access ou une feuille `.xls` porte souvent le nom du client,
    et la conversion l'écrit tel quel : en nom d'onglet, en clé de premier
    niveau d'un JSON, en nom de fichier. Contrairement au classeur d'où
    une formule le référence, ici RIEN ne le résout — c'est donc de la
    donnée, et le tolérer aurait laissé le nom du client dans la copie
    tout en faisant refuser le fichier au filet.

    Il passe par la MÊME table que les cellules : la feuille et les
    valeurs qui la nomment reçoivent le même mot.
    """
    from script.data.external_file import nouveau_mot

    if table is None or not feuille.nom:
        return feuille.nom
    return nouveau_mot(str(feuille.nom), table, options["vivier"])


def convertir(chemin, destination, cible, options, feuilles=None, table=None):
    """Écrire la copie dans un AUTRE format. Rend la liste des fichiers.

    Quand la cible ne peut pas porter la forme de la source — un classeur
    de plusieurs feuilles vers csv, une base Access vers csv — la
    destination est un RÉPERTOIRE, un fichier par feuille. Le menu le dit
    avant d'écrire.
    """
    if cible not in CIBLES_CONVERSION:
        raise ErreurMoteur("format_inconnu", cible)
    if feuilles is None:
        feuilles = _preparer(chemin, options)[1]
    retenues = [
        f
        for f in feuilles
        if not options.get("feuilles") or f.nom in options["feuilles"]
    ]
    if not retenues:
        raise ErreurMoteur("aucune_feuille", "")

    noms = {
        f.nom: _nom_de_feuille_anonyme(f, table, options) for f in retenues
    }
    if cible == "xlsx":
        return _convertir_vers_xlsx(destination, retenues, noms)
    if cible == "csv":
        if len(retenues) == 1:
            return _ecrire_csv(destination, retenues[0], options)
        return _convertir_vers_repertoire(destination, retenues, options, noms)
    if cible == "json":
        return _convertir_vers_json(destination, retenues, noms)
    return _convertir_vers_xml(destination, retenues)


def _convertir_vers_xlsx(destination, feuilles, noms=None):
    from openpyxl import Workbook

    classeur = Workbook()
    classeur.remove(classeur.active)
    pris = set()
    for feuille in feuilles:
        brut = (noms or {}).get(feuille.nom, feuille.nom)
        onglet = classeur.create_sheet(nom_de_fichier_sur(brut, pris, 31))
        for ligne in feuille.lignes:
            onglet.append([_valeur_pour_xlsx(v) for v in ligne])
    _ecrire_atomique(destination, classeur.save)
    return [destination]


def _valeur_pour_xlsx(valeur):
    """La valeur telle qu'un classeur la porte, sans passer par du texte.

    `valeur_hors_tableur` est écrite pour csv, json et xml, trois formats
    sans types : elle rend une date en chaîne ISO. openpyxl, lui, porte
    nativement `datetime`, `int`, `float` et `bool` — convertir en texte
    faisait perdre le type, et la copie portait du texte là où une date
    était attendue, donc ne se réimportait plus comme telle.
    """
    if valeur is None or isinstance(valeur, (bool, int, float, str)):
        return valeur
    texte = getattr(valeur, "text", None)
    if isinstance(texte, str):
        return texte
    if isinstance(valeur, (bytes, bytearray)):
        return None
    if isinstance(valeur, (dict, list)):
        # Un conteneur imbriqué d'un JSON : openpyxl lève sur une cellule
        # qu'il ne sait pas porter, et l'erreur ressortait brute sous
        # « format non reconnu ». Son contenu est déjà anonymisé.
        return json.dumps(valeur, ensure_ascii=False)
    return valeur


def _preparer_repertoire(destination):
    if os.path.isdir(destination) and os.listdir(destination):
        raise ErreurMoteur("repertoire_non_vide", destination)
    os.makedirs(destination, mode=0o700, exist_ok=True)


def _convertir_vers_repertoire(destination, feuilles, options, noms=None):
    _preparer_repertoire(destination)
    pris = set()
    ecrits = []
    for feuille in feuilles:
        nom = nom_de_fichier_sur(
            (noms or {}).get(feuille.nom, feuille.nom), pris
        )
        cible = os.path.join(destination, f"{nom}.csv")
        ecrits.extend(_ecrire_csv(cible, feuille, options))
    return ecrits


def _noms_neutres(feuille):
    """Des noms de colonne quand la feuille n'a PAS d'en-tête.

    Sans eux, la première ligne de données prenait la place des noms de
    clé : en clair dans la copie, et perdue comme donnée puisqu'un objet
    ne porte pas sa clé deux fois.
    """
    largeur = max((len(l) for l in feuille.lignes), default=0)
    return ["colonne_%d" % (i + 1) for i in range(largeur)]


def _cles_distinctes(etiquettes):
    """Une clé d'objet par colonne, toutes DISTINCTES.

    Un objet JSON écrase la clé qu'il répète : deux colonnes au même
    en-tête, ou deux en-têtes vides, n'en laissaient qu'une dans la copie
    et la donnée de l'autre disparaissait sans qu'une ligne ne le dise.
    Un tableur laisse les deux faire — l'en-tête n'y est qu'une ligne.

    Une étiquette vide n'est pas None : elle porte la chaîne vide, que le
    repli sur « cN » ne couvrait pas. Le test porte sur l'étiquette
    dépouillée, la clé garde l'étiquette telle quelle.
    """
    sortie = []
    pris = set()
    for rang, valeur in enumerate(etiquettes, start=1):
        nom = "" if valeur is None else str(valeur)
        if not nom.strip():
            nom = "c%d" % rang
        candidat, suffixe = nom, 2
        while candidat in pris:
            candidat = "%s_%d" % (nom, suffixe)
            suffixe += 1
        pris.add(candidat)
        sortie.append(candidat)
    return sortie


def _convertir_vers_json(destination, feuilles, noms=None):
    sortie = {}
    # La clé de premier niveau nomme la feuille, et passe par la même
    # distinction que ses colonnes : les deux autres graveurs réservent
    # déjà leurs noms d'onglet et de fichier, celui-ci écrasait.
    cles = _cles_distinctes([(noms or {}).get(f.nom, f.nom) for f in feuilles])
    for cle, feuille in zip(cles, feuilles):
        etiquettes = _cles_distinctes(
            feuille.etiquettes or _noms_neutres(feuille)
        )
        enregistrements = []
        for ligne in corps(feuille):
            enregistrements.append(
                {
                    etiquettes[i]: valeur_hors_tableur(v)
                    for i, v in enumerate(ligne)
                    if i < len(etiquettes)
                }
            )
        sortie[cle] = enregistrements

    def ecrivain(cible):
        with open(cible, "w", encoding="utf-8") as fh:
            json.dump(sortie, fh, ensure_ascii=False, indent=1)

    _ecrire_atomique(destination, ecrivain)
    return [destination]


def _nom_de_balise_sur(etiquette, rang):
    """Une étiquette rendue légale comme nom d'élément XML.

    XML interdit à un nom de commencer par un chiffre : une colonne
    intitulée « 2024 » produisait `<2024>`, écrit sans broncher, annoncé
    comme écrit, et refusé par tout analyseur — y compris celui du dépôt.
    """
    brut = re.sub(r"[^A-Za-z0-9_.-]", "_", str(etiquette or "")).strip("_")
    if not brut:
        return f"c{rang}"
    if not re.match(r"[A-Za-z_]", brut[0]):
        return f"c{rang}_{brut}"
    return brut


def _convertir_vers_xml(destination, feuilles):
    if len(feuilles) > 1:
        raise ErreurMoteur("format_inconnu", "xml")
    feuille = feuilles[0]
    racine = ET.Element("table")
    etiquettes = [
        _nom_de_balise_sur(v, i)
        for i, v in enumerate(
            feuille.etiquettes or _noms_neutres(feuille), start=1
        )
    ]
    for ligne in corps(feuille):
        noeud = ET.SubElement(racine, "ligne")
        for index, valeur in enumerate(ligne):
            if index >= len(etiquettes):
                continue
            champ = ET.SubElement(noeud, etiquettes[index])
            rendu = valeur_hors_tableur(valeur)
            champ.text = "" if rendu is None else str(rendu)

    def ecrivain(cible):
        ET.ElementTree(racine).write(
            cible, encoding="utf-8", xml_declaration=True
        )

    _ecrire_atomique(destination, ecrivain)
    return [destination]
